"""Estimate Availability windows from printed Attendance sheets.

For every active member (open Enrollment row), reads the Time-In /
Time-Out values on the last three months of Attendance sheets, takes the
per-weekday envelope (earliest Time-In, latest Time-Out) and writes it to
the currently-open Availability row. One transaction; --dry-run rolls
back. A report CSV lists old/new windows and flags.

See docs/superpowers/specs/2026-09-15-attendance-availability-backfill-design.md
"""
import argparse
import csv
import datetime
import os
import re
import shutil
import sys
from pathlib import Path

# Make `from monthly_schedule import ...` work when invoked as a script.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from monthly_schedule.attendance_envelope import (  # noqa: E402
    DEFAULT_MIN_SAMPLES,
    collect_samples,
    estimate_windows,
    hhmm,
    normalize_time,
    parse_sheet_filename,
    window_flags,
)

DEFAULT_SHEETS_ROOT = r"\\Dell-NJ02\Desktop\Backup TP System\data"

# Rows 3..33 of the Attendance sheet: A=day number, B=weekday, C=Time-In,
# D=Time-Out. Row 2 is the header.
_FIRST_DATA_ROW = 3
_LAST_DATA_ROW = 33

_DAY_NAMES = {1: "Mon", 2: "Tue", 3: "Wed", 4: "Thu", 5: "Fri",
              6: "Sat", 7: "Sun"}


# ---------------------------------------------------------------------------
# months
# ---------------------------------------------------------------------------

def _default_months(today):
    """The three 'YYYY-MM' months ending with `today`'s month."""
    year, month = today.year, today.month
    out = []
    for _ in range(3):
        out.append(f"{year:04d}-{month:02d}")
        month -= 1
        if month == 0:
            month, year = 12, year - 1
    return list(reversed(out))


def _parse_months(text):
    months = [m.strip() for m in text.split(",") if m.strip()]
    for m in months:
        if not re.fullmatch(r"\d{4}-\d{2}", m):
            raise ValueError(f"bad month {m!r}; expected YYYY-MM")
    return months


def _today():
    """Wrapped so tests can pin the date."""
    return datetime.date.today()


# ---------------------------------------------------------------------------
# sheets
# ---------------------------------------------------------------------------

def read_month_sheets(root, month, on_error):
    """Yield (center_id, iso_weekday, in_min, out_min) for every dated row
    of every Attendance sheet under <root>/<YYYY>/<MM>/Attendance/.

    Rows with no visit (blank times) are yielded with None so the caller's
    collect_samples can skip them. Unreadable workbooks and a missing
    month folder are reported through on_error(str) and skipped.
    """
    import openpyxl

    year_s, mon_s = month.split("-")
    folder = Path(root) / year_s / mon_s / "Attendance"
    if not folder.is_dir():
        on_error(f"month folder not found for {month}: {folder}")
        return
    year, mon = int(year_s), int(mon_s)
    for path in sorted(folder.iterdir()):
        parsed = parse_sheet_filename(path.name)
        if parsed is None:
            continue
        center_id, _name, _month = parsed
        try:
            wb = openpyxl.load_workbook(
                path, read_only=True, data_only=True, keep_vba=False)
        except Exception as exc:  # openpyxl raises many types
            on_error(f"could not open {path.name}: {exc}")
            continue
        try:
            ws = wb["Attendance"] if "Attendance" in wb.sheetnames \
                else wb.worksheets[-1]
            for row in ws.iter_rows(min_row=_FIRST_DATA_ROW,
                                    max_row=_LAST_DATA_ROW, max_col=4,
                                    values_only=True):
                day_no, _dow, time_in, time_out = row
                if not isinstance(day_no, (int, float)):
                    continue
                try:
                    date = datetime.date(year, mon, int(day_no))
                except ValueError:
                    continue   # 29..31 on a shorter month
                yield (center_id, date.isoweekday(),
                       normalize_time(time_in), normalize_time(time_out))
        finally:
            wb.close()


# ---------------------------------------------------------------------------
# args
# ---------------------------------------------------------------------------

def _parse_args(argv):
    p = argparse.ArgumentParser(
        description="Estimate Availability windows from Attendance sheets.",
    )
    p.add_argument("--db", required=True,
                   help="Path to the Access .accdb file.")
    p.add_argument("--sheets-root", default=DEFAULT_SHEETS_ROOT,
                   help="Folder holding <YYYY>/<MM>/Attendance/ "
                        f"(default {DEFAULT_SHEETS_ROOT}).")
    p.add_argument("--months", default=None,
                   help="Comma-separated YYYY-MM list. Default: the three "
                        "months ending this month.")
    p.add_argument("--csv-out", default=None,
                   help="Directory for the report CSV. "
                        "Default: the DB's directory.")
    p.add_argument("--min-samples", type=int, default=DEFAULT_MIN_SAMPLES,
                   help="Weekdays with fewer samples widen to the "
                        f"member-wide envelope (default {DEFAULT_MIN_SAMPLES}).")
    p.add_argument("--dry-run", action="store_true",
                   help="Report only; roll back DB writes; no backup copy.")
    p.add_argument("--quiet", action="store_true",
                   help="Suppress per-member stdout; print only the summary.")
    return p.parse_args(argv)


# ---------------------------------------------------------------------------
# database
# ---------------------------------------------------------------------------

_ACTIVE_MEMBERS_QUERY = (
    "SELECT c.[Center ID], c.[Last Name], c.[First Name], "
    "c.[Health Plan], c.[HHA] "
    "FROM [Contacts] c "
    "WHERE c.[Center ID] IN "
    "(SELECT e.[Center ID] FROM [Enrollment] e WHERE e.[end_date] IS NULL) "
    "ORDER BY c.[Center ID]"
)
_OPERATING_DAYS_QUERY = (
    "SELECT [Day Of Week], [closing_time] FROM [OperatingDays] ORDER BY [ID]"
)
_AVAIL_OPEN_QUERY = (
    "SELECT [ID], [avail_start], [avail_end] FROM [Availability] "
    "WHERE [Center ID] = ? AND [Day Of Week] = ? "
    "AND [effective_end_date] IS NULL"
)
_AVAIL_UPDATE = (
    "UPDATE [Availability] SET [avail_start] = ?, [avail_end] = ?, "
    "[Notes] = ? WHERE [ID] = ?"
)
_AVAIL_INSERT = (
    "INSERT INTO [Availability] "
    "([Center ID], [effective_start_date], [effective_end_date], "
    "[Day Of Week], [avail_start], [avail_end], [Notes]) "
    "VALUES (?, ?, NULL, ?, ?, ?, ?)"
)

_FALLBACK_CLOSING_MIN = 16 * 60   # rules.py latest_time_out when no row


def _minutes_to_time(minutes):
    return datetime.time(minutes // 60, minutes % 60)


def _time_to_minutes(value):
    """Access stores time-of-day as DATETIME with a 1899-12-30 date part;
    pyodbc returns datetime.datetime. Accept datetime.time too."""
    t = value.time() if hasattr(value, "time") else value
    return t.hour * 60 + t.minute


def _fetch_active_members(cur):
    """Active = has an Enrollment row with end_date IS NULL."""
    cur.execute(_ACTIVE_MEMBERS_QUERY)
    members = []
    for cid, last, first, plan, hha in cur.fetchall():
        if cid is None:
            continue
        members.append({
            "center_id": int(cid),
            "last_name": (last or "").strip(),
            "first_name": (first or "").strip(),
            "health_plan": (plan or "").strip(),
            "hha": (hha or "").strip(),
        })
    return members


def _fetch_closing_times(cur):
    """{iso_weekday: closing minutes}; the largest ID wins per weekday
    (matches CenterCalendar); weekdays without a row fall back to 16:00."""
    closing = {d: _FALLBACK_CLOSING_MIN for d in range(1, 8)}
    cur.execute(_OPERATING_DAYS_QUERY)
    for dow, closing_time in cur.fetchall():
        if dow is None or closing_time is None:
            continue
        closing[int(dow)] = _time_to_minutes(closing_time)
    return closing


def _fetch_open_window(cur, center_id, day):
    """(row_id, (start_min, end_min)) for the open row, or None."""
    cur.execute(_AVAIL_OPEN_QUERY, str(center_id), day)
    row = cur.fetchone()
    if row is None:
        return None
    row_id, start, end = row
    return int(row_id), (_time_to_minutes(start), _time_to_minutes(end))


def _upsert_window(cur, center_id, day, start_min, end_min, notes, today):
    """Write the window to the open row (UPDATE) or a new row (INSERT).
    Returns the previous (start_min, end_min) or None if there was no row."""
    found = _fetch_open_window(cur, center_id, day)
    start_t, end_t = _minutes_to_time(start_min), _minutes_to_time(end_min)
    if found is None:
        cur.execute(_AVAIL_INSERT, str(center_id), today, day,
                    start_t, end_t, notes)
        return None
    row_id, old = found
    cur.execute(_AVAIL_UPDATE, start_t, end_t, notes, row_id)
    return old


def _build_connection_string(db_path):
    return (
        "DRIVER={Microsoft Access Driver (*.mdb, *.accdb)};"
        f"DBQ={db_path};"
    )


def main(argv=None):  # filled in by Task 8
    raise NotImplementedError


if __name__ == "__main__":
    sys.exit(main())
