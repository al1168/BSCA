# -*- coding: utf-8 -*-
"""Per-member Bowery activity log workbook, filled from the bundled
Excel template.

Generated only when the GUI "Program name" setting contains "bowery".
The day grid mirrors the timesheet's eligibility decisions via the
`status` key carried on each timesheet row ("attended" / "absent" /
"ineligible"): attended days get 2-4 reproducibly-random checkmarks
among the activities offered that weekday, absent days get a merged
"Absent" label, and every other day is grayed out.
"""

import calendar
import datetime as dt
import os
import random
import sys
from copy import copy

from openpyxl import load_workbook
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.styles import Alignment, Font, PatternFill

from monthly_schedule.billing_workbook import parse_flex_date
from monthly_schedule.month_dates import get_month_dates

CHECK_CHAR = "ü"                   # renders as a checkmark in Wingdings
CHECK_FONT = Font(name="Wingdings", size=11)
# Set explicitly on every check cell — a few template cells (e.g. C37)
# carry no alignment of their own.
CHECK_ALIGNMENT = Alignment(horizontal="center", vertical="center")
ABSENT_TEXT = "Absent"
ABSENT_FONT = Font(name="Calibri", size=11)
ABSENT_ALIGNMENT = Alignment(horizontal="center", vertical="center")
GRAY_RGB = "FFD9D9D9"
GRAY_FILL = PatternFill(fill_type="solid",
                        start_color=GRAY_RGB, end_color=GRAY_RGB)

ACTIVITY_COUNT = 14
ACTIVITY_FIRST_COL = 3                  # C; A1 -> C ... A14 -> P
ACTIVITY_LAST_COL = ACTIVITY_FIRST_COL + ACTIVITY_COUNT - 1
DATE_FIRST_ROW = 8
TEMPLATE_DATE_ROWS = 30                 # the shipped template is a 30-day June
DATE_NUMFMT = "m/d/yyyy;@"


def _asset_base():
    """Root for bundled assets: sys._MEIPASS when frozen by PyInstaller
    (must match the spec's datas dest), else the repo root."""
    base = getattr(sys, "_MEIPASS", None)
    if base is None:
        base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    return base


def default_template_path():
    """Path to the bundled Bowery template."""
    return os.path.join(_asset_base(), "assets", "templates",
                        "Bowery activity log template.xlsx")


def is_bowery_program(program_name):
    """True iff the Settings "Program name" contains 'bowery' (any
    casing). Kept for compatibility; new code uses activity_program."""
    return "bowery" in str(program_name or "").lower()


def activity_program(program_name):
    """Which activity log the Settings "Program name" selects:
    'bowery', 'cathay', or None (no logs). Each program maps to its own
    template and layout. Bowery wins if a name somehow contains both,
    preserving the pre-Cathay behavior."""
    name = str(program_name or "").lower()
    if "bowery" in name:
        return "bowery"
    if "cathay" in name:
        return "cathay"
    return None


def parse_frequency(raw):
    """'1.2.3.4.5' -> frozenset({1..5}) (1=Mon..7=Sun). Blank/None ->
    empty set (activity never offered); non-digit or out-of-range
    tokens are ignored."""
    days = set()
    for token in str(raw or "").split("."):
        token = token.strip()
        if token.isdigit() and 1 <= int(token) <= 7:
            days.add(int(token))
    return frozenset(days)


def activity_log_filename(center_id, year, month):
    return (f"({center_id}) {calendar.month_abbr[month]} {year} "
            f"Activity log.xlsx")


def activity_log_subdir(year, month, health_plan):
    """All-members runs group logs by insurance plan:
    'Activity Logs 2026-06/HF'. Normalization mirrors
    all_members_subdir (blank/None -> '_NoPlan')."""
    plan = (str(health_plan).strip().upper()
            if health_plan else "") or "_NoPlan"
    return os.path.join(f"Activity Logs {year:04d}-{month:02d}", plan)


def _indexes_offered(day, activities, activity_count):
    """Sorted 1-based activity indexes offered on `day`'s weekday
    (activity A<i>); ids outside A1..A<activity_count> are ignored."""
    weekday = day.isoweekday()
    indexes = []
    for i in range(1, activity_count + 1):
        info = activities.get(f"A{i}")
        if info and weekday in parse_frequency(info.get("frequency")):
            indexes.append(i)
    return indexes


def pick_activity_indexes(center_id, day, activities,
                          activity_count=ACTIVITY_COUNT):
    """2-4 reproducibly-random 1-based activity indexes to check for an
    attended day, restricted to activities offered that weekday. Seeded
    per member and per day (string seed = stable across Python runs) so
    re-runs give identical files and one day's picks never depend on
    another day's status. The loop shape must stay a plain
    range(1, n+1) scan — reordering would silently reshuffle
    historical outputs."""
    available = _indexes_offered(day, activities, activity_count)
    if not available:
        return []
    rng = random.Random(f"activity:{center_id}:{day.isoformat()}")
    k = min(rng.randint(2, 4), len(available))
    return sorted(rng.sample(available, k))


def pick_activity_columns(center_id, day, activities):
    """Bowery layout: activity A<i> lives in column 2+i."""
    return [ACTIVITY_FIRST_COL - 1 + i
            for i in pick_activity_indexes(center_id, day, activities)]


def latest_auth_member_id(ctx, year, month):
    """The insurance Member ID from the latest authorization active in
    the month (same latest-wins rule the timesheet's Auth label uses).
    '' when no authorization covers any day."""
    for day in reversed(get_month_dates(year, month)):
        auth = ctx.active_authorization(day)
        if auth is not None:
            return str(auth.get("member_id") or "").strip()
    return ""


def build_activity_log(member, rows, activities, year, month, *,
                       dob=None, member_id="", template_path=None):
    """Fill a copy of the template for one member and return the
    Workbook (caller saves via save_activity_log).

    `rows` are the timesheet rows from build_rows (only `date` and
    `status` are read); days missing from `rows` (partial-range runs)
    are grayed out like ineligible days."""
    wb = load_workbook(template_path or default_template_path())
    ws = wb["Sheet1"]

    # The generator paints every day explicitly; leftover conditional-
    # formatting rules in the template (e.g. weekend highlighting)
    # would fight those fills, so drop them all.
    ws.conditional_formatting = ConditionalFormattingList()

    # Header block.
    ws["B3"] = int(member["center_id"])
    ws["E3"] = f"{member['last_name']}, {member['first_name']}"
    parsed_dob = parse_flex_date(dob)
    if parsed_dob is not None:
        ws["H3"] = parsed_dob
    else:
        ws["H3"] = None
    plan_cell = ws["K3"]
    plan_cell.value = str(member.get("health_plan") or "").strip()
    plan_cell.number_format = "General"  # template carries a stray format
    ws["O3"] = member_id or None

    # Activity name headers (blank when the id is missing from the DB).
    for i in range(1, ACTIVITY_COUNT + 1):
        info = activities.get(f"A{i}", {})
        col = ACTIVITY_FIRST_COL - 1 + i
        # .value assignment (not ws.cell(value=...)) so a missing
        # activity blanks the template's literal "A3" header.
        ws.cell(row=6, column=col).value = info.get("name") or None
        ws.cell(row=7, column=col).value = info.get("c_name") or None

    # Drop any sample checkmarks shipped in the template's date grid.
    for r in range(DATE_FIRST_ROW, DATE_FIRST_ROW + TEMPLATE_DATE_ROWS):
        for c in range(ACTIVITY_FIRST_COL, ACTIVITY_LAST_COL + 1):
            ws.cell(row=r, column=c).value = None

    # Resize the 30-row June grid to this month's length.
    n_days = calendar.monthrange(year, month)[1]
    last_row = DATE_FIRST_ROW + n_days - 1
    if n_days < TEMPLATE_DATE_ROWS:
        ws.delete_rows(last_row + 1, TEMPLATE_DATE_ROWS - n_days)
    elif n_days > TEMPLATE_DATE_ROWS:
        # insert_rows adds unstyled cells; clone the style of the last
        # clean date row so borders/formats survive.
        template_row = DATE_FIRST_ROW + TEMPLATE_DATE_ROWS - 1
        for extra in range(TEMPLATE_DATE_ROWS, n_days):
            new_row = DATE_FIRST_ROW + extra
            ws.insert_rows(new_row)
            for col in range(1, ACTIVITY_LAST_COL + 1):
                source = ws.cell(row=template_row, column=col)
                ws.cell(row=new_row, column=col)._style = copy(source._style)

    # Rewrite every date/weekday cell for the target month.
    for d in range(1, n_days + 1):
        r = DATE_FIRST_ROW + d - 1
        date_cell = ws.cell(row=r, column=1,
                            value=dt.datetime(year, month, d))
        date_cell.number_format = DATE_NUMFMT
        ws.cell(row=r, column=2, value=f'=TEXT(A{r},"ddd")')

    # Fill the day rows from the timesheet's eligibility decisions.
    status_by_day = {row["date"].day: row.get("status") for row in rows}
    for d in range(1, n_days + 1):
        r = DATE_FIRST_ROW + d - 1
        status = status_by_day.get(d)
        if status == "attended":
            day = dt.date(year, month, d)
            for col in pick_activity_columns(member["center_id"], day,
                                             activities):
                cell = ws.cell(row=r, column=col, value=CHECK_CHAR)
                cell.font = CHECK_FONT
                cell.alignment = CHECK_ALIGNMENT
        elif status == "absent":
            ws.merge_cells(start_row=r, start_column=ACTIVITY_FIRST_COL,
                           end_row=r, end_column=ACTIVITY_LAST_COL)
            cell = ws.cell(row=r, column=ACTIVITY_FIRST_COL,
                           value=ABSENT_TEXT)
            cell.font = ABSENT_FONT
            cell.alignment = ABSENT_ALIGNMENT
        else:  # ineligible, or outside a partial-range run
            # Full-row gray (date/week included), matching the
            # template's original weekend highlighting.
            for col in range(1, ACTIVITY_LAST_COL + 1):
                ws.cell(row=r, column=col).fill = GRAY_FILL
    return wb


def save_activity_log(wb, out_dir, center_id, year, month,
                      on_fallback=None):
    """Save into `out_dir`; when the primary name is locked (open in
    Excel), retry `_1`..`_9` suffixes; `on_fallback(primary, actual)` is
    called once when a fallback is used. Raises PermissionError only
    when every candidate is locked. Returns the path written."""
    stem, ext = os.path.splitext(
        activity_log_filename(center_id, year, month))
    primary = os.path.join(out_dir, f"{stem}{ext}")
    candidates = [primary] + [
        os.path.join(out_dir, f"{stem}_{n}{ext}") for n in range(1, 10)
    ]
    last_error = None
    for path in candidates:
        try:
            wb.save(path)
        except PermissionError as exc:
            last_error = exc
            continue
        if path != primary and on_fallback is not None:
            on_fallback(primary, path)
        return path
    raise last_error
