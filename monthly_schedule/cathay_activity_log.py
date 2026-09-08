# -*- coding: utf-8 -*-
"""Per-member Cathay activity log workbook, filled from the bundled
macro-enabled template (saved as plain .xlsx — macros are dropped).

Selected when the GUI "Program name" contains "cathay"; the Bowery
program has its own template and builder (activity_log_workbook.py).

Layout is transposed vs Bowery: activities A1..A24 run down rows 9..32
(labels in column A, rewritten from the DB's Activities table), days of
the month 1..31 run across columns B..AF. Only days the timesheet
scheduled (row status "attended") get 2-4 reproducibly-random '✓'
checks; absent/unauthorized days are simply left empty — no gray
fills, no "Absent" labels.
"""

import calendar
import datetime as dt
import os

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

from monthly_schedule.activity_log_workbook import (
    _asset_base,
    pick_activity_indexes,
)

SHEET_NAME = "AttndActivityLog"
CHECK_CHAR = "✓"                   # '✓'
CHECK_FONT = Font(name="Calibri", size=10, bold=True)
CHECK_ALIGNMENT = Alignment(horizontal="center", vertical="center")

ACTIVITY_COUNT = 24
ACTIVITY_FIRST_ROW = 9                  # A<i> -> row 8 + i
DAY_FIRST_COL = 2                       # day d -> column 1 + d (B..AF)
DAY_LAST_COL = DAY_FIRST_COL + 31 - 1   # AF


def cathay_template_path():
    return os.path.join(_asset_base(), "assets", "templates",
                        "Cathay activity log template.xlsm")


def build_cathay_activity_log(member, rows, activities, year, month, *,
                              auth_weekdays=frozenset(),
                              template_path=None):
    """Fill a copy of the Cathay template for one member and return the
    Workbook (caller saves via save_activity_log, producing .xlsx).

    `rows` are the timesheet rows from build_rows (only `date` and
    `status` are read); any day not marked "attended" — absent,
    unauthorized, or missing from a partial-range run — is left
    completely empty."""
    # keep_vba=False drops the template's vbaProject so the saved .xlsx
    # is a clean macro-free workbook.
    wb = load_workbook(template_path or cathay_template_path(),
                       keep_vba=False)
    ws = wb[SHEET_NAME]

    # Header line: '{NAME}       ID: {CENTER_ID}   (AUTH_DAYS)'.
    # Token-only replacement keeps the template's spacing/font/merge.
    text = str(ws["B6"].value or "")
    text = text.replace(
        "{NAME}", f"{member['last_name']}, {member['first_name']}")
    text = text.replace("{CENTER_ID}", str(member["center_id"]))
    auth_text = ".".join(str(d) for d in sorted(auth_weekdays))
    if auth_text:
        text = text.replace("(AUTH_DAYS)", f"({auth_text})")
    else:
        text = text.replace("(AUTH_DAYS)", "").rstrip()
    ws["B6"] = text

    # Date range: '{MM/DD/YYYY} - {MM/DD/YYYY}' -> '4/1/2026 - 4/30/2026'
    # (no zero-padding, matching the sample).
    n_days = calendar.monthrange(year, month)[1]
    range_text = str(ws["B7"].value or "")
    for day_num in (1, n_days):
        range_text = range_text.replace(
            "{MM/DD/YYYY}", f"{month}/{day_num}/{year}", 1)
    ws["B7"] = range_text

    # Activity labels, database wins: '{A1:ACTIVITY_NAME}' -> 'A1:Cards'.
    # Blank the row's label when the A_ID is missing from the DB.
    for i in range(1, ACTIVITY_COUNT + 1):
        name = (activities.get(f"A{i}") or {}).get("name")
        ws.cell(row=ACTIVITY_FIRST_ROW - 1 + i, column=1).value = (
            f"A{i}:{name}" if name else None)

    # The shipped template is a filled sample — clear every check in
    # the full 31-column grid before writing fresh ones.
    for r in range(ACTIVITY_FIRST_ROW,
                   ACTIVITY_FIRST_ROW + ACTIVITY_COUNT):
        for c in range(DAY_FIRST_COL, DAY_LAST_COL + 1):
            ws.cell(row=r, column=c).value = None

    # Checks on attended days only (same seed as Bowery, so re-runs
    # reproduce identical files). Everything else stays untouched,
    # including day columns beyond this month's length.
    status_by_day = {row["date"].day: row.get("status") for row in rows}
    for d in range(1, n_days + 1):
        if status_by_day.get(d) != "attended":
            continue
        day = dt.date(year, month, d)
        col = DAY_FIRST_COL - 1 + d
        for i in pick_activity_indexes(member["center_id"], day,
                                       activities,
                                       activity_count=ACTIVITY_COUNT):
            cell = ws.cell(row=ACTIVITY_FIRST_ROW - 1 + i, column=col,
                           value=CHECK_CHAR)
            cell.font = CHECK_FONT
            cell.alignment = CHECK_ALIGNMENT
    return wb
