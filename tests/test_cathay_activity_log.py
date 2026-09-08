# -*- coding: utf-8 -*-
"""Tests for the Cathay activity log workbook (transposed layout:
activities as rows A1..A24, days 1..31 as columns)."""

import calendar
import datetime as dt
import os

import pytest
from openpyxl import load_workbook

from monthly_schedule.activity_log_workbook import (
    activity_program,
    is_bowery_program,
    pick_activity_columns,
    pick_activity_indexes,
    save_activity_log,
)
from monthly_schedule.cathay_activity_log import (
    CHECK_CHAR,
    build_cathay_activity_log,
    cathay_template_path,
)

MEMBER = {
    "center_id": 25347,
    "last_name": "Chan",
    "first_name": "Mei",
    "health_plan": "AE",
}

# A1..A24 like the Cathay Activities table; A20 weekend-only, A24 never
# offered (blank frequency).
ACTIVITIES = {}
for _i in range(1, 25):
    ACTIVITIES[f"A{_i}"] = {
        "name": f"Act {_i}",
        "c_name": f"活动{_i}",
        "frequency": "1.2.3.4.5.6.7",
    }
ACTIVITIES["A20"] = {"name": "Yoga", "c_name": "瑜伽", "frequency": "6.7"}
ACTIVITIES["A24"] = {"name": "Group", "c_name": "小组", "frequency": ""}

ACTIVITIES_14 = {k: v for k, v in ACTIVITIES.items()
                 if int(k[1:]) <= 14}


def make_rows(year, month, attended_weekdays=(1, 2, 3),
              absent_days=(10,)):
    n_days = calendar.monthrange(year, month)[1]
    rows = []
    for d in range(1, n_days + 1):
        day = dt.date(year, month, d)
        if d in absent_days:
            status = "absent"
        elif day.isoweekday() in attended_weekdays:
            status = "attended"
        else:
            status = "ineligible"
        rows.append({"date": day, "day": "x", "status": status})
    return rows


# --- program resolver -------------------------------------------------


def test_activity_program_cathay_variations():
    assert activity_program("Cathay") == "cathay"
    assert activity_program("CATHAY ADC") == "cathay"
    assert activity_program(" cathay adult daycare ") == "cathay"


def test_activity_program_bowery():
    assert activity_program("Bowery SADC") == "bowery"


def test_activity_program_none():
    assert activity_program("") is None
    assert activity_program(None) is None
    assert activity_program("Chinatown") is None


def test_activity_program_bowery_wins_collision():
    assert activity_program("Bowery and Cathay") == "bowery"


def test_is_bowery_program_still_works():
    assert is_bowery_program("Bowery SADC")
    assert not is_bowery_program("Cathay")


# --- generalized picker -----------------------------------------------


def test_pick_activity_indexes_two_to_four():
    day = dt.date(2026, 4, 1)  # Wednesday
    idx = pick_activity_indexes(25347, day, ACTIVITIES, activity_count=24)
    assert 2 <= len(idx) <= 4
    assert idx == sorted(idx)
    assert all(1 <= i <= 24 for i in idx)


def test_pick_activity_indexes_reaches_high_indexes():
    # With 24 activities available, indexes above 14 must be reachable
    # on some day of the month.
    seen = set()
    for d in range(1, 31):
        seen.update(pick_activity_indexes(
            25347, dt.date(2026, 4, d), ACTIVITIES, activity_count=24))
    assert any(i > 14 for i in seen)


def test_pick_activity_indexes_respects_frequency():
    for d in range(1, 6):  # Wed..Sun of April 2026 starts Wed; use Mon-Fri
        day = dt.date(2026, 4, d)
        if day.isoweekday() > 5:
            continue
        idx = pick_activity_indexes(25347, day, ACTIVITIES,
                                    activity_count=24)
        assert 20 not in idx  # weekend-only
        assert 24 not in idx  # never offered


def test_pick_activity_columns_parity_with_indexes():
    # The Bowery column picker must remain byte-identical to the
    # generalized index picker shifted by 2 — this pins that the
    # refactor cannot reshuffle historical Bowery outputs.
    for d in (1, 2, 5, 14, 28):
        day = dt.date(2026, 6, d)
        assert pick_activity_columns(25347, day, ACTIVITIES_14) == [
            i + 2 for i in pick_activity_indexes(25347, day, ACTIVITIES_14)
        ]


# --- template path ----------------------------------------------------


def test_cathay_template_path_unfrozen():
    path = cathay_template_path()
    assert path.endswith(os.path.join(
        "assets", "templates", "Cathay activity log template.xlsm"))
    assert os.path.exists(path)


def test_cathay_template_path_frozen(monkeypatch):
    import sys
    monkeypatch.setattr(sys, "_MEIPASS", r"C:\frozen\base", raising=False)
    assert cathay_template_path().startswith(r"C:\frozen\base")


# --- full build (April 2026, 30 days) ---------------------------------


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    rows = make_rows(2026, 4)
    wb = build_cathay_activity_log(MEMBER, rows, ACTIVITIES, 2026, 4,
                                   auth_weekdays={1, 3, 5})
    out = tmp_path_factory.mktemp("cathay") / "log.xlsx"
    wb.save(str(out))
    return load_workbook(str(out))


def test_b6_tokens_replaced(built):
    ws = built["AttndActivityLog"]
    assert ws["B6"].value == "Chan, Mei       ID: 25347   (1.3.5)"


def test_b6_empty_auth_drops_parenthetical():
    rows = make_rows(2026, 4)
    wb = build_cathay_activity_log(MEMBER, rows, ACTIVITIES, 2026, 4,
                                   auth_weekdays=frozenset())
    text = wb["AttndActivityLog"]["B6"].value
    assert "AUTH_DAYS" not in text
    assert "()" not in text
    assert text.rstrip() == text


def test_b7_date_range(built):
    ws = built["AttndActivityLog"]
    assert ws["B7"].value == "4/1/2026 - 4/30/2026"


def test_activity_labels_from_db(built):
    ws = built["AttndActivityLog"]
    assert ws["A9"].value == "A1:Act 1"
    assert ws["A28"].value == "A20:Yoga"
    assert ws["A32"].value == "A24:Group"


def test_missing_activity_id_blanks_label():
    activities = dict(ACTIVITIES)
    del activities["A3"]
    rows = make_rows(2026, 4)
    ws = build_cathay_activity_log(MEMBER, rows, activities, 2026, 4)[
        "AttndActivityLog"]
    assert ws["A11"].value in (None, "")


def test_no_leftover_tokens(built):
    ws = built["AttndActivityLog"]
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert "{" not in cell.value, cell.coordinate


def test_sample_checks_cleared():
    # The shipped template holds 56 sample checks (first at C9); a
    # build with no attended days must leave the whole grid empty.
    rows = make_rows(2026, 4, attended_weekdays=(), absent_days=())
    ws = build_cathay_activity_log(MEMBER, rows, ACTIVITIES, 2026, 4)[
        "AttndActivityLog"]
    assert ws["C9"].value in (None, "")
    for r in range(9, 33):
        for c in range(2, 33):
            assert ws.cell(row=r, column=c).value in (None, ""), (r, c)


def test_attended_day_checks(built):
    ws = built["AttndActivityLog"]
    # 2026-04-01 is a Wednesday (attended) -> column B=2.
    checks = [ws.cell(row=r, column=2) for r in range(9, 33)
              if ws.cell(row=r, column=2).value == CHECK_CHAR]
    assert 2 <= len(checks) <= 4
    for cell in checks:
        assert cell.font.name == "Calibri"
        assert cell.font.size == 10
        assert cell.font.bold
        assert cell.alignment.horizontal == "center"
        assert cell.alignment.vertical == "center"


def test_absent_day_completely_empty(built):
    ws = built["AttndActivityLog"]
    # Day 10 is absent -> column K=11: no checks, no "Absent" text.
    for r in range(9, 33):
        assert ws.cell(row=r, column=11).value in (None, "")


def test_no_absent_text_or_fills_anywhere(built):
    ws = built["AttndActivityLog"]
    for row in ws.iter_rows(min_row=9, max_row=32, min_col=2, max_col=32):
        for cell in row:
            assert cell.value in (None, "", CHECK_CHAR)
            # theme-colored white background from the template is fine;
            # the generator must not introduce rgb fills like the
            # Bowery gray.
            if cell.fill.fill_type == "solid":
                assert cell.fill.start_color.type != "rgb" or \
                    cell.fill.start_color.rgb in (None, "00000000")


def test_ineligible_day_empty(built):
    ws = built["AttndActivityLog"]
    # 2026-04-04 is a Saturday -> ineligible -> column E=5.
    for r in range(9, 33):
        assert ws.cell(row=r, column=5).value in (None, "")


def test_day_beyond_month_untouched(built):
    ws = built["AttndActivityLog"]
    # April has 30 days; column AF=32 keeps its header "31", no checks.
    assert str(ws.cell(row=8, column=32).value) == "31"
    for r in range(9, 33):
        assert ws.cell(row=r, column=32).value in (None, "")


def test_days_missing_from_rows_empty():
    rows = make_rows(2026, 4)[:5]  # partial-range run
    ws = build_cathay_activity_log(MEMBER, rows, ACTIVITIES, 2026, 4)[
        "AttndActivityLog"]
    for c in range(7, 33):  # days 6..31
        for r in range(9, 33):
            assert ws.cell(row=r, column=c).value in (None, "")


def test_weekend_only_activity_never_checked_weekdays(built):
    ws = built["AttndActivityLog"]
    for d in range(1, 31):
        if dt.date(2026, 4, d).isoweekday() <= 5:
            assert ws.cell(row=28, column=1 + d).value != CHECK_CHAR


def test_merges_preserved(built):
    merged = {str(r) for r in built["AttndActivityLog"].merged_cells.ranges}
    assert {"A1:AF4", "B6:P6", "B7:AF7"} <= merged


def test_reproducible_and_day_independent():
    base = make_rows(2026, 4)
    changed = make_rows(2026, 4, absent_days=(2, 10))
    a = build_cathay_activity_log(MEMBER, base, ACTIVITIES, 2026, 4)[
        "AttndActivityLog"]
    b = build_cathay_activity_log(MEMBER, base, ACTIVITIES, 2026, 4)[
        "AttndActivityLog"]
    c = build_cathay_activity_log(MEMBER, changed, ACTIVITIES, 2026, 4)[
        "AttndActivityLog"]
    coords = lambda ws: {(r, col) for r in range(9, 33)
                         for col in range(2, 33)
                         if ws.cell(row=r, column=col).value == CHECK_CHAR}
    assert coords(a) == coords(b)
    assert coords(a)
    # Flipping day 2 to absent leaves day 1's column (B=2) unchanged.
    col_b = lambda ws: [ws.cell(row=r, column=2).value for r in range(9, 33)]
    assert col_b(a) == col_b(c)


def test_save_as_xlsx(tmp_path):
    rows = make_rows(2026, 4)
    wb = build_cathay_activity_log(MEMBER, rows, ACTIVITIES, 2026, 4)
    path = save_activity_log(wb, str(tmp_path), 25347, 2026, 4)
    assert os.path.basename(path) == "(25347) Apr 2026 Activity log.xlsx"
    reloaded = load_workbook(path)  # valid macro-free xlsx
    assert "AttndActivityLog" in reloaded.sheetnames
