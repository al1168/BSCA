from datetime import date, datetime

from openpyxl import load_workbook

from monthly_schedule.workbook import build_workbook, COMPANY_NAME

MEMBER = {
    "center_id": 24010,
    "last_name": "Cheng",
    "first_name": "Lizhu",
    "health_plan": "HOF",
}

ROWS = [
    {
        "date": date(2026, 5, 1), "day": "Fri",
        "pickup": "08:05", "arrival": "08:15", "time_in": "08:17",
        "time_out": "12:13", "departure": "12:15", "dropoff": "12:25",
    },
    {
        "date": date(2026, 5, 2), "day": "Sat",
        "pickup": "", "arrival": "", "time_in": "",
        "time_out": "", "departure": "", "dropoff": "",
    },
]


def test_side_by_side_workbook_structure(tmp_path):
    out = tmp_path / "sched.xlsx"
    build_workbook(MEMBER, ROWS, str(out))
    assert out.exists()

    wb = load_workbook(str(out))
    ws = wb["Schedule"]

    # per-block header: left block at col A, right block at col F
    assert ws["A1"].value == COMPANY_NAME
    assert ws["F1"].value == COMPANY_NAME
    assert ws["A2"].value == "MLTC: Elderplan Homefirst"
    assert ws["F2"].value == "MLTC: Elderplan Homefirst"
    for cell in ("A3", "F3"):
        v = ws[cell].value
        assert "ID: 24010" in v
        assert "Name: Cheng, Lizhu" in v

    merged = {str(rng) for rng in ws.merged_cells.ranges}
    for rng in ("A1:D1", "A2:D2", "A3:D3",
                "F1:K1", "F2:K2", "F3:K3"):
        assert rng in merged

    # table-header row is row 4 for both blocks
    assert [ws.cell(row=4, column=c).value
            for c in range(1, 5)] == [
        "Date", "Day", "Time-In", "Time-Out"
    ]
    assert [ws.cell(row=4, column=c).value
            for c in range(6, 12)] == [
        "Date", "Day", "Pick-Up Time", "Arrival Time",
        "Departure Time", "Drop-Off Time",
    ]

    # first data row (row 5), left and right blocks aligned
    assert ws.cell(row=5, column=1).value == datetime(2026, 5, 1)
    assert ws.cell(row=5, column=1).number_format == "m/d/yyyy"
    assert ws.cell(row=5, column=3).value == "08:17"   # Time-In
    assert ws.cell(row=5, column=4).value == "12:13"   # Time-Out
    assert ws.cell(row=5, column=6).value == datetime(2026, 5, 1)
    assert ws.cell(row=5, column=6).number_format == "m/d/yyyy"
    assert ws.cell(row=5, column=8).value == "08:05"   # Pick-Up
    assert ws.cell(row=5, column=9).value == "08:15"   # Arrival
    assert ws.cell(row=5, column=10).value == "12:15"  # Departure
    assert ws.cell(row=5, column=11).value == "12:25"  # Drop-Off

    # ineligible row (row 6) blank in both blocks
    assert ws.cell(row=6, column=3).value in (None, "")
    assert ws.cell(row=6, column=8).value in (None, "")

    # exactly one COLUMN break at column 5; no row breaks
    assert ws.col_breaks.count == 1
    assert list(ws.col_breaks.brk)[0].id == 5
    assert ws.row_breaks.count == 0

    # print area spans to column K, through the footer (row 11)
    assert ws.print_area is not None
    assert "K" in ws.print_area
    assert "K11" in ws.print_area.replace("$", "")

    # spacer column fixed width (unchanged)
    assert ws.column_dimensions["E"].width == 3

    # page setup: not fit-to-width, fit each page to one page tall
    assert ws.page_setup.fitToWidth == 0
    assert ws.page_setup.fitToHeight == 1
    # tight margins so a full month fits one page at ~100% (no shrink
    # that would drop the thin grid lines in Excel's print scaling)
    assert ws.page_margins.top <= 0.35
    assert ws.page_margins.bottom <= 0.35

    # --- footer: bold caption merged per block ---
    # ROWS has 2 entries -> last_data_row = 6, caption row 8. Two blank
    # rows (9, 10) of space, then the signature/date line at row 11.
    assert ws["A8"].value == "Attendance Sheet"
    assert ws["F8"].value == "Transportation Sheet"
    assert "A8:D8" in merged
    assert "F8:K8" in merged
    # the two spacer rows between caption and signature carry no labels
    assert ws.cell(row=9, column=1).value is None
    assert ws.cell(row=10, column=1).value is None

    # --- signature/date row 11: each label shares ONE merged cell with
    #     its extended write-on line, bottom-bordered across the span ---
    assert ws.cell(row=11, column=1).value == "Signature:"
    assert ws.cell(row=11, column=4).value == "Date:"
    assert ws.cell(row=11, column=6).value == "Signature:"
    assert ws.cell(row=11, column=9).value == "Date:"
    # the label cell AND its rule cells all carry the bottom rule
    for c in (1, 2, 3, 4, 6, 7, 8, 9, 10, 11):
        assert ws.cell(row=11, column=c).border.bottom.style == "thin"
    # label + line is one merged cell per field, in both footers; the
    # left Signature takes A:C (like the transport F:H) so its line is
    # the long one, leaving the Date label + line on column D alone
    for rng in ("A11:C11", "F11:H11", "I11:K11"):
        assert rng in merged
    assert "A11:B11" not in merged
    assert "C11:D11" not in merged

    # --- horizontal print centering on ---
    assert ws.print_options.horizontalCentered is True

    # --- footer text did NOT bloat a data column ---
    # (autosize is bounded to the data region; col F is the right
    # Date col ~13.5 and is not a rule column)
    assert ws.column_dimensions["F"].width < 20

    # --- Attendance (left) column proportions: Date readable, Day
    #     narrow (like the transport Day col), Time-In/Out widened ---
    assert ws.column_dimensions["A"].width >= 14          # Date
    assert ws.column_dimensions["B"].width < 10           # Day, narrow
    assert ws.column_dimensions["C"].width >= 22          # Time-In, wide
    assert ws.column_dimensions["D"].width >= 22          # Time-Out, wide

    # --- guaranteed Signature/Date line lengths: the merged write-on
    #     spans (A:C sig, D date) stay long enough without forcing the
    #     Day column wide ---
    assert (ws.column_dimensions["A"].width
            + ws.column_dimensions["B"].width
            + ws.column_dimensions["C"].width) >= 22 - 1e-6   # left sig
    assert ws.column_dimensions["D"].width >= 14 - 1e-6       # left date
    assert (ws.column_dimensions["G"].width
            + ws.column_dimensions["H"].width) >= 22 - 1e-6
    assert (ws.column_dimensions["J"].width
            + ws.column_dimensions["K"].width) >= 14 - 1e-6

    # --- both Day columns are narrow (left now matches right) ---
    assert ws.column_dimensions["G"].width < 16

    # --- larger font + taller rows applied ---
    assert ws.cell(row=5, column=3).font.size == 12   # a data cell
    assert ws.cell(row=4, column=1).font.size == 12   # a header cell
    assert ws.row_dimensions[5].height == 18          # a data row

    # --- all header lines are bold (company, MLTC, ID/Name) ---
    assert ws.cell(row=1, column=1).font.bold is True
    assert ws.cell(row=2, column=1).font.bold is True   # MLTC
    assert ws.cell(row=3, column=1).font.bold is True   # ID / Name


def test_header_shows_authorized_weekdays(tmp_path):
    out = tmp_path / "sched.xlsx"
    build_workbook(MEMBER, ROWS, str(out), auth_weekdays={1, 5, 7})

    wb = load_workbook(str(out))
    ws = wb["Schedule"]
    merged = {str(rng) for rng in ws.merged_cells.ranges}

    # Company spans the full width on row 1 — auth no longer sits there.
    assert ws["A1"].value == COMPANY_NAME
    for rng in ("A1:D1", "F1:K1", "A2:D2", "F2:K2"):
        assert rng in merged

    # Auth days move down onto the ID/Name row (row 3), right-aligned,
    # directly on top of the table, dot-separated.
    assert ws.cell(row=3, column=4).value == "Auth: 1.5.7"    # D3 (left)
    assert ws.cell(row=3, column=10).value == "Auth: 1.5.7"   # J3 (right)
    assert ws.cell(row=3, column=4).alignment.horizontal == "right"
    assert ws.cell(row=3, column=10).alignment.horizontal == "right"
    assert ws.cell(row=3, column=4).font.bold is True
    assert ws.cell(row=3, column=10).font.bold is True
    # ID/Name keeps the rest of the row; auth gets the rightmost col(s).
    assert "A3:C3" in merged      # left ID/Name
    assert "F3:I3" in merged      # right ID/Name
    assert "J3:K3" in merged      # right auth spans 2 columns
    assert ws.cell(row=3, column=1).value.startswith("ID:")


def test_header_without_auth_weekdays_keeps_full_width(tmp_path):
    out = tmp_path / "sched.xlsx"
    build_workbook(MEMBER, ROWS, str(out))  # no auth_weekdays
    wb = load_workbook(str(out))
    ws = wb["Schedule"]
    merged = {str(rng) for rng in ws.merged_cells.ranges}
    # Every header row stays one merged span per table (back-compat).
    for rng in ("A1:D1", "F1:K1", "A3:D3", "F3:K3"):
        assert rng in merged
