import calendar
import csv
import datetime
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from scripts import backfill_availability_from_attendance as mod


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def _t(hour, minute=0):
    """Time-of-day as Access/pyodbc returns it."""
    return datetime.datetime(1899, 12, 30, hour, minute)


def _write_attendance_sheet(folder, center_id, name, month, days):
    """Create '<folder>/(id).Name Attendance YYYY-MM.xlsm' with one row per
    calendar day; `days` maps day-number -> (time_in, time_out) as
    datetime.time (12-hour, no AM/PM, like the real sheets)."""
    folder.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.active.title = "Template"
    ws = wb.create_sheet("Attendance")
    ws["A1"] = f"Name: {name}(1.2.3)   Month, 2026  ID: {center_id}  MLTC: X"
    ws.append(["Date", "Day", "Time-In", "Time-Out"])
    year, mon = (int(p) for p in month.split("-"))
    for d in range(1, calendar.monthrange(year, mon)[1] + 1):
        dow = datetime.date(year, mon, d).strftime("%a")
        ti, to = days.get(d, (None, None))
        ws.append([d, dow, ti, to])
    path = folder / f"({center_id}).{name} Attendance {month}.xlsm"
    wb.save(path)   # xlsx bytes under an .xlsm name; the reader keys on the name
    return path


# ---------------------------------------------------------------------------
# _default_months / _parse_months
# ---------------------------------------------------------------------------

def test_default_months_three_ending_this_month():
    assert mod._default_months(datetime.date(2026, 9, 15)) == [
        "2026-07", "2026-08", "2026-09"]


def test_default_months_wraps_year():
    assert mod._default_months(datetime.date(2026, 1, 3)) == [
        "2025-11", "2025-12", "2026-01"]


def test_parse_months_flag():
    assert mod._parse_months("2026-07, 2026-08,2026-09") == [
        "2026-07", "2026-08", "2026-09"]
    with pytest.raises(ValueError):
        mod._parse_months("July 2026")


# ---------------------------------------------------------------------------
# read_month_sheets — yields (center_id, iso_weekday, in_min, out_min)
# ---------------------------------------------------------------------------

def test_read_month_sheets_yields_normalized_rows(tmp_path):
    att = tmp_path / "2026" / "08" / "Attendance"
    _write_attendance_sheet(att, 1001, "Zhang, Mingli", "2026-08", {
        3: (datetime.time(8, 21), datetime.time(12, 23)),   # Mon
        4: (datetime.time(9, 41), datetime.time(2, 0)),     # Tue, 02:00 -> 14:00
    })
    (att / "Daily Sign-in-out 2026-08.xlsm").write_bytes(b"")  # ignored by name
    errors = []
    rows = [r for r in mod.read_month_sheets(tmp_path, "2026-08", errors.append)
            if r[2] is not None or r[3] is not None]
    assert rows == [
        (1001, 1, 8 * 60 + 21, 12 * 60 + 23),
        (1001, 2, 9 * 60 + 41, 14 * 60),
    ]
    assert errors == []


def test_read_month_sheets_reports_unreadable_file(tmp_path):
    att = tmp_path / "2026" / "07" / "Attendance"
    att.mkdir(parents=True)
    (att / "(5).Bad, File Attendance 2026-07.xlsm").write_bytes(b"not a workbook")
    errors = []
    rows = list(mod.read_month_sheets(tmp_path, "2026-07", errors.append))
    assert rows == []
    assert len(errors) == 1 and "(5).Bad, File" in errors[0]


def test_read_month_sheets_missing_folder_is_an_error(tmp_path):
    errors = []
    rows = list(mod.read_month_sheets(tmp_path, "2026-06", errors.append))
    assert rows == []
    assert len(errors) == 1 and "2026-06" in errors[0]
