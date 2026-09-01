import new_monthly_schedule as cli

from datetime import date


def _fake_enrollments(cid=24010):
    return [{"id": 1, "center_id": cid,
             "start_date": date(2026, 1, 1), "end_date": None}]


def _fake_authorizations(cid=24010, auth_days="1.3.4.5"):
    return [{"id": 1, "center_id": cid,
             "auth_start": date(2026, 1, 1),
             "auth_end": date(2026, 12, 31),
             "effective_start": date(2026, 1, 1),
             "effective_end": date(2026, 12, 31),
             "auth_days": auth_days}]


FAKE_MEMBER = {
    "center_id": 24010,
    "last_name": "Cheng",
    "first_name": "Lizhu",
    "health_plan": "HOF",
    "address": "1 Main St, NY",
    "long_lat": None,
}

FAKE_MEMBER_2 = {
    "center_id": 24011,
    "last_name": "Smith",
    "first_name": "John",
    "health_plan": "HOF",
    "address": "2 Main St, NY",
    "long_lat": None,
}


import pytest


@pytest.fixture(autouse=True)
def _stub_travel(monkeypatch):
    """Neutralize travel + new DB lookups for legacy tests.

    Injects --api-key K into every parse_args call so individual tests
    don't have to thread the flag through every cli.main() invocation.
    """
    _real_parse_args = cli.parse_args

    def _parse_args_with_key(argv):
        if "--api-key" not in argv:
            argv = list(argv) + ["--api-key", "K"]
        return _real_parse_args(argv)

    monkeypatch.setattr(cli, "parse_args", _parse_args_with_key)
    monkeypatch.setattr(cli, "load_cache", lambda path: {})
    monkeypatch.setattr(cli, "save_cache", lambda path, cache: None)
    monkeypatch.setattr(cli, "load_time_cache", lambda path: {})
    monkeypatch.setattr(cli, "save_time_cache", lambda path, cache: None)
    monkeypatch.setattr(
        cli, "resolve_travel_minutes",
        lambda member, api_key, cache: 10,
    )
    monkeypatch.setattr(
        cli, "get_enrollments",
        lambda cid, db: _fake_enrollments(cid),
    )
    monkeypatch.setattr(
        cli, "get_authorizations",
        lambda cid, db: _fake_authorizations(cid),
    )
    monkeypatch.setattr(cli, "get_absences", lambda cid, db: [])
    monkeypatch.setattr(cli, "get_availability", lambda cid, db: [])
    monkeypatch.setattr(cli, "get_one_offs", lambda cid, db: [])


def test_no_member_returns_2(monkeypatch, tmp_path, capsys):
    monkeypatch.setattr(cli, "get_member", lambda cid, db: None)
    rc = cli.main(
        ["--center-id", "999", "--year", "2026", "--month", "5",
         "--output-path", str(tmp_path)]
    )
    assert rc == 2
    err = capsys.readouterr().err
    assert "ID 999: lookup — not found in database" in err


def test_writes_workbook(monkeypatch, tmp_path):
    monkeypatch.setattr(cli, "get_member", lambda cid, db: FAKE_MEMBER)
    rc = cli.main(
        ["--center-id", "24010", "--year", "2026", "--month", "5",
         "--output-path", str(tmp_path)]
    )
    assert rc == 0
    assert (tmp_path / "Schedule_24010_2026-05.xlsx").exists()


def test_invalid_month_rejected(monkeypatch):
    monkeypatch.setattr(cli, "get_member", lambda cid, db: FAKE_MEMBER)
    import pytest
    with pytest.raises(SystemExit):
        cli.main(
            ["--center-id", "24010", "--year", "2026", "--month", "13"]
        )


def test_missing_db_returns_1(monkeypatch, capsys):
    def _raise(cid, db):
        raise FileNotFoundError("Database not found: X")
    monkeypatch.setattr(cli, "get_member", _raise)
    rc = cli.main(["--center-id", "24010", "--year", "2026", "--month", "5"])
    assert rc == 1
    assert "Database not found" in capsys.readouterr().err


def test_driver_error_returns_1(monkeypatch, capsys):
    def _raise(cid, db):
        raise RuntimeError("Could not open the Access database. ...")
    monkeypatch.setattr(cli, "get_member", _raise)
    rc = cli.main(["--center-id", "24010", "--year", "2026", "--month", "5"])
    assert rc == 1
    assert "Could not open the Access database" in capsys.readouterr().err


def test_main_reads_sys_argv_when_argv_none(monkeypatch, tmp_path):
    """Exercises the argv is None branch (real CLI invocation path)."""
    monkeypatch.setattr(cli, "get_member", lambda cid, db: FAKE_MEMBER)
    monkeypatch.setattr(cli, "build_workbook", lambda *a, **k: None)
    monkeypatch.setattr(
        cli.sys, "argv",
        ["new_monthly_schedule.py", "--center-id", "24010",
         "--year", "2026", "--month", "5",
         "--output-path", str(tmp_path)],
    )
    assert cli.main() == 0


def test_parse_center_ids_basic():
    assert cli.parse_center_ids("24010,24011") == [24010, 24011]


def test_parse_center_ids_trims_dedupes_drops_blanks_keeps_order():
    assert cli.parse_center_ids("24011, 24010 ,,24011") == [24011, 24010]


def test_parse_center_ids_empty():
    assert cli.parse_center_ids("") == []
    assert cli.parse_center_ids(" , , ") == []


def test_parse_center_ids_non_int_raises():
    with pytest.raises(ValueError):
        cli.parse_center_ids("24010,abc")


import os


def test_schedule_filename():
    assert cli.schedule_filename(24010, 2026, 5) == \
        "Schedule_24010_2026-05.xlsx"


def test_resolve_output_dir_non_plan_is_base():
    assert cli.resolve_output_dir(".", None, 2026, 5) == "."
    assert cli.resolve_output_dir("out", None, 2026, 5) == "out"


def test_resolve_output_dir_plan_nests_uppercased_subdir():
    assert cli.resolve_output_dir("out", "hof", 2026, 5) == \
        os.path.join("out", "HOF_2026-05")


def test_format_summary_all_success_headline_only():
    s = cli.format_summary("Wrote", 2, 2, "2026-05", "out", [])
    assert s == "Wrote 2 of 2 member(s) for 2026-05 into out; 0 failed."
    assert "Failures:" not in s


def test_format_summary_with_failures_lists_them():
    failures = [
        cli.Failure(24099, "", "lookup", "not found in database", None),
        cli.Failure(24100, "Smith, John", "write",
                    "PermissionError — denied", None),
    ]
    s = cli.format_summary(
        "Wrote", 18, 20, "plan HOF 2026-05", "out/HOF_2026-05", failures
    )
    lines = s.split("\n")
    assert lines[0] == (
        "Wrote 18 of 20 member(s) for plan HOF 2026-05 "
        "into out/HOF_2026-05; 2 failed."
    )
    assert lines[1] == "Failures:"
    assert lines[2] == "  - ID 24099: lookup — not found in database"
    assert lines[3] == (
        "  - ID 24100 (Smith, John): write — "
        "PermissionError — denied"
    )


def test_parse_args_requires_a_selector():
    with pytest.raises(SystemExit):
        cli.parse_args(["--year", "2026", "--month", "5"])


def test_parse_args_rejects_two_selectors():
    with pytest.raises(SystemExit):
        cli.parse_args(
            ["--center-id", "1", "--plan", "HOF",
             "--year", "2026", "--month", "5"]
        )


def test_parse_args_accepts_each_selector():
    a = cli.parse_args(
        ["--center-ids", "1,2", "--year", "2026", "--month", "5"]
    )
    assert a.center_ids == "1,2"
    assert a.center_id is None and a.plan is None
    b = cli.parse_args(
        ["--plan", "HOF", "--year", "2026", "--month", "5"]
    )
    assert b.plan == "HOF"
    assert b.output_path == "."  # new default: base directory


def test_single_writes_into_output_dir(monkeypatch, tmp_path):
    monkeypatch.setattr(cli, "get_member", lambda cid, db: FAKE_MEMBER)
    rc = cli.main(
        ["--center-id", "24010", "--year", "2026", "--month", "5",
         "--output-path", str(tmp_path)]
    )
    assert rc == 0
    assert (tmp_path / "Schedule_24010_2026-05.xlsx").exists()


def test_center_ids_batch_one_not_found(monkeypatch, tmp_path, capsys):
    def fake_get_member(cid, db):
        return FAKE_MEMBER if cid == 24010 else None
    monkeypatch.setattr(cli, "get_member", fake_get_member)
    rc = cli.main(
        ["--center-ids", "24010,24099", "--year", "2026",
         "--month", "5", "--output-path", str(tmp_path)]
    )
    assert rc == 2
    assert (tmp_path / "Schedule_24010_2026-05.xlsx").exists()
    err = capsys.readouterr().err
    assert "Wrote 1 of 2 member(s) for 2026-05" in err
    assert "Failures:" in err
    assert "ID 24099: lookup — not found in database" in err


def test_plan_writes_into_subdir(monkeypatch, tmp_path, capsys):
    monkeypatch.setattr(
        cli, "get_members_by_plan",
        lambda code, db: [FAKE_MEMBER, FAKE_MEMBER_2],
    )
    rc = cli.main(
        ["--plan", "hof", "--year", "2026", "--month", "5",
         "--output-path", str(tmp_path)]
    )
    assert rc == 0
    sub = tmp_path / "HOF_2026-05"
    assert (sub / "Schedule_24010_2026-05.xlsx").exists()
    assert (sub / "Schedule_24011_2026-05.xlsx").exists()
    assert "for plan HOF 2026-05" in capsys.readouterr().err


def test_plan_no_members(monkeypatch, tmp_path, capsys):
    monkeypatch.setattr(
        cli, "get_members_by_plan", lambda code, db: []
    )
    rc = cli.main(
        ["--plan", "HOF", "--year", "2026", "--month", "5",
         "--output-path", str(tmp_path)]
    )
    assert rc == 2
    assert "No members found for plan HOF" in capsys.readouterr().err
    assert not (tmp_path / "HOF_2026-05").exists()


def test_write_failure_reported_in_summary(
        monkeypatch, tmp_path, capsys):
    monkeypatch.setattr(cli, "get_member", lambda cid, db: FAKE_MEMBER)

    def boom(member, rows, path, **kwargs):
        raise PermissionError("denied")
    monkeypatch.setattr(cli, "build_workbook", boom)
    rc = cli.main(
        ["--center-id", "24010", "--year", "2026", "--month", "5",
         "--output-path", str(tmp_path)]
    )
    assert rc == 2
    err = capsys.readouterr().err
    assert "Failures:" in err
    assert "ID 24010 (Cheng, Lizhu): write — PermissionError — denied" \
        in err


def test_travel_minutes_applied_to_pickup_and_dropoff(
        monkeypatch, tmp_path):
    monkeypatch.setattr(cli, "get_member", lambda cid, db: FAKE_MEMBER)
    monkeypatch.setattr(
        cli, "resolve_travel_minutes",
        lambda member, api_key, cache: 7,
    )
    captured = {}

    def _capture(member, rows, path, **kwargs):
        captured["rows"] = rows
    monkeypatch.setattr(cli, "build_workbook", _capture)

    rc = cli.main(
        ["--center-id", "24010", "--year", "2026", "--month", "5",
         "--output-path", str(tmp_path)]
    )
    assert rc == 0

    def to_min(hhmm):
        h, m = hhmm.split(":")
        return int(h) * 60 + int(m)

    from monthly_schedule.rules import SCHEDULE_RULES
    buf_lo, buf_hi = SCHEDULE_RULES["Default"]["travel_buffer_min"]
    travel_minutes = 7

    seen = False
    for row in captured["rows"]:
        if row["arrival"] == "":
            continue
        lead = to_min(row["arrival"]) - to_min(row["pickup"])
        trail = to_min(row["dropoff"]) - to_min(row["departure"])
        assert travel_minutes + buf_lo <= lead <= travel_minutes + buf_hi
        assert travel_minutes + buf_lo <= trail <= travel_minutes + buf_hi
        seen = True
    assert seen  # at least one eligible day was checked


def test_travel_failure_skips_member_in_summary(
        monkeypatch, tmp_path, capsys):
    monkeypatch.setattr(cli, "get_member", lambda cid, db: FAKE_MEMBER)

    def boom(member, api_key, cache):
        raise cli.TravelError("geocode", "ZERO_RESULTS for 'x'")
    monkeypatch.setattr(cli, "resolve_travel_minutes", boom)
    rc = cli.main(
        ["--center-id", "24010", "--year", "2026", "--month", "5",
         "--output-path", str(tmp_path)]
    )
    assert rc == 2
    err = capsys.readouterr().err
    assert "Failures:" in err
    assert ("ID 24010 (Cheng, Lizhu): geocode — "
            "ZERO_RESULTS for 'x'") in err
    assert not (tmp_path / "Schedule_24010_2026-05.xlsx").exists()


def test_no_enrollment_yields_failure(monkeypatch, tmp_path, capsys):
    monkeypatch.setattr(cli, "get_member", lambda cid, db: FAKE_MEMBER)
    monkeypatch.setattr(cli, "get_enrollments", lambda cid, db: [])
    rc = cli.main(
        ["--center-id", "24010", "--year", "2026", "--month", "5",
         "--output-path", str(tmp_path)]
    )
    assert rc == 2
    err = capsys.readouterr().err
    assert "eligibility — not enrolled during this month" in err


def test_no_authorization_yields_failure(monkeypatch, tmp_path, capsys):
    monkeypatch.setattr(cli, "get_member", lambda cid, db: FAKE_MEMBER)
    monkeypatch.setattr(cli, "get_authorizations", lambda cid, db: [])
    rc = cli.main(
        ["--center-id", "24010", "--year", "2026", "--month", "5",
         "--output-path", str(tmp_path)]
    )
    assert rc == 2
    err = capsys.readouterr().err
    assert "eligibility — no active authorization for this month" in err


# ---------------------------------------------------------------------------
# write_skipped_members_csv
# ---------------------------------------------------------------------------

from new_monthly_schedule import write_skipped_members_csv, Failure


def test_write_skipped_members_csv_empty_returns_none(tmp_path):
    """Empty failures list returns None and writes nothing."""
    result = write_skipped_members_csv(
        [], str(tmp_path), today=date(2026, 6, 4)
    )
    assert result is None
    assert list(tmp_path.iterdir()) == []


def test_write_skipped_members_csv_filename_uses_today(tmp_path):
    """The returned path is `skipped_members_<today>.csv` in out_dir."""
    failures = [
        Failure(24010, "Cheng, Lizhu", "lookup", "not found", None),
    ]
    result = write_skipped_members_csv(
        failures, str(tmp_path), today=date(2026, 6, 4)
    )
    expected_path = tmp_path / "skipped_members_2026-06-04.csv"
    assert result == str(expected_path)
    assert expected_path.exists()
    lines = expected_path.read_text(encoding="utf-8").splitlines()
    assert lines[0] == "center_id,name,stage,reason,day"


def test_write_skipped_members_csv_includes_all_stages(tmp_path):
    """All failure stages appear in the CSV. The `day` column is the ISO
    date only for one_off_conflict rows and empty for everything else."""
    failures = [
        Failure(24010, "Cheng, Lizhu", "lookup", "not found in database", None),
        Failure(24011, "Smith, John", "eligibility", "not enrolled", None),
        Failure(24012, "Jones, Mary", "geocode", "ZERO_RESULTS for 'x'", None),
        Failure(24013, "Park, Eun", "one_off_conflict",
                "one-off on auth day", date(2026, 6, 5)),
    ]
    result = write_skipped_members_csv(
        failures, str(tmp_path), today=date(2026, 6, 4)
    )
    expected_path = tmp_path / "skipped_members_2026-06-04.csv"
    assert result == str(expected_path)
    content = expected_path.read_text(encoding="utf-8")
    lines = content.splitlines()
    assert lines[0] == "center_id,name,stage,reason,day"
    assert len(lines) == 5  # header + 4 failures
    assert lines[1] == '24010,"Cheng, Lizhu",lookup,not found in database,'
    assert lines[2] == '24011,"Smith, John",eligibility,not enrolled,'
    assert lines[3] == '24012,"Jones, Mary",geocode,ZERO_RESULTS for \'x\','
    assert lines[4] == (
        '24013,"Park, Eun",one_off_conflict,one-off on auth day,2026-06-05'
    )


def test_write_skipped_members_csv_falls_back_when_primary_locked(tmp_path):
    """When the primary file can't be opened for writing (typically because
    the CSV from an earlier run today is open in Excel), the report is
    written to `skipped_members_<date>_1.csv` instead and `on_fallback`
    receives both paths. A directory squatting on the primary name raises
    PermissionError on Windows, same as Excel's lock."""
    (tmp_path / "skipped_members_2026-06-04.csv").mkdir()
    failures = [
        Failure(24010, "Cheng, Lizhu", "lookup", "not found", None),
    ]
    calls = []
    result = write_skipped_members_csv(
        failures, str(tmp_path), today=date(2026, 6, 4),
        on_fallback=lambda primary, actual: calls.append((primary, actual)),
    )
    expected_path = tmp_path / "skipped_members_2026-06-04_1.csv"
    assert result == str(expected_path)
    lines = expected_path.read_text(encoding="utf-8").splitlines()
    assert lines[0] == "center_id,name,stage,reason,day"
    assert len(lines) == 2
    assert calls == [
        (str(tmp_path / "skipped_members_2026-06-04.csv"),
         str(expected_path)),
    ]


def test_write_skipped_members_csv_tries_next_fallback_name(tmp_path):
    """If the first fallback name is also locked, the next one is used."""
    (tmp_path / "skipped_members_2026-06-04.csv").mkdir()
    (tmp_path / "skipped_members_2026-06-04_1.csv").mkdir()
    failures = [
        Failure(24010, "Cheng, Lizhu", "lookup", "not found", None),
    ]
    result = write_skipped_members_csv(
        failures, str(tmp_path), today=date(2026, 6, 4)
    )
    expected_path = tmp_path / "skipped_members_2026-06-04_2.csv"
    assert result == str(expected_path)
    assert expected_path.exists()


def test_write_skipped_members_csv_raises_when_all_names_locked(tmp_path):
    """Primary plus every fallback name locked -> PermissionError."""
    (tmp_path / "skipped_members_2026-06-04.csv").mkdir()
    for n in range(1, 10):
        (tmp_path / f"skipped_members_2026-06-04_{n}.csv").mkdir()
    failures = [
        Failure(24010, "Cheng, Lizhu", "lookup", "not found", None),
    ]
    with pytest.raises(PermissionError):
        write_skipped_members_csv(
            failures, str(tmp_path), today=date(2026, 6, 4)
        )


def test_cli_warns_and_falls_back_when_skipped_csv_locked(
        monkeypatch, tmp_path, capsys):
    """When the primary skipped-members CSV is locked (e.g. open in
    Excel), main() warns on stderr and writes the report under the `_1`
    fallback name instead of crashing."""
    monkeypatch.setattr(cli, "get_member", lambda cid, db: None)
    today = date.today().isoformat()
    (tmp_path / f"skipped_members_{today}.csv").mkdir()
    rc = cli.main(
        ["--center-id", "24010", "--year", "2026", "--month", "5",
         "--output-path", str(tmp_path)]
    )
    assert rc == 2
    err = capsys.readouterr().err
    fallback = tmp_path / f"skipped_members_{today}_1.csv"
    assert fallback.exists()
    assert "open in Excel" in err
    assert f"skipped_members_{today}_1.csv" in err
    assert "Failures:" in err


def test_cli_warns_when_skipped_csv_unwritable(
        monkeypatch, tmp_path, capsys):
    """When every candidate CSV name is locked, main() warns on stderr
    and still prints the failure summary instead of crashing."""
    monkeypatch.setattr(cli, "get_member", lambda cid, db: None)
    today = date.today().isoformat()
    (tmp_path / f"skipped_members_{today}.csv").mkdir()
    for n in range(1, 10):
        (tmp_path / f"skipped_members_{today}_{n}.csv").mkdir()
    rc = cli.main(
        ["--center-id", "24010", "--year", "2026", "--month", "5",
         "--output-path", str(tmp_path)]
    )
    assert rc == 2
    err = capsys.readouterr().err
    assert "could not write the skipped members report" in err
    assert "Failures:" in err


def test_write_skipped_members_csv_no_callback_when_primary_writable(tmp_path):
    """`on_fallback` is not called when the primary name is used."""
    failures = [
        Failure(24010, "Cheng, Lizhu", "lookup", "not found", None),
    ]
    calls = []
    result = write_skipped_members_csv(
        failures, str(tmp_path), today=date(2026, 6, 4),
        on_fallback=lambda primary, actual: calls.append((primary, actual)),
    )
    assert result == str(tmp_path / "skipped_members_2026-06-04.csv")
    assert calls == []


# ---------------------------------------------------------------------------
# Custom day range
# ---------------------------------------------------------------------------


def test_range_suffix_added_to_schedule_filename():
    assert cli.schedule_filename(24010, 2026, 5, 1, 19) == \
        "Schedule_24010_2026-05_01-19.xlsx"


def test_no_range_omits_filename_suffix():
    assert cli.schedule_filename(24010, 2026, 5) == \
        "Schedule_24010_2026-05.xlsx"


def test_debug_filename_with_range():
    assert (cli.debug_filename(24010, 2026, 5, 1, 19)
            == "Debug_24010_2026-05_01-19.csv")


def test_debug_filename_without_range():
    assert cli.debug_filename(24010, 2026, 5) == "Debug_24010_2026-05.csv"


def test_debug_subdir():
    assert cli.debug_subdir(2026, 5) == "Debug 2026-05"


def test_all_members_subdir_single_folder():
    # Default: everyone shares one month-named folder, plan-independent.
    assert cli.all_members_subdir(2026, 6, "HOF", False) == "June_2026_Timesheets"
    assert cli.all_members_subdir(2026, 6, None, False) == "June_2026_Timesheets"


def test_all_members_subdir_per_plan():
    # Opt-in: one folder per MLTC, matching the old behavior.
    assert cli.all_members_subdir(2026, 6, "hof", True) == "HOF_2026-06"
    assert cli.all_members_subdir(2026, 6, None, True) == "_NoPlan_2026-06"
    assert cli.all_members_subdir(2026, 6, "   ", True) == "_NoPlan_2026-06"


def test_cli_range_writes_to_range_suffixed_file(monkeypatch, tmp_path):
    """--start-day/--end-day produce a schedule with the suffixed name."""
    monkeypatch.setattr(cli, "get_member", lambda cid, db: FAKE_MEMBER)
    rc = cli.main(
        ["--center-id", "24010", "--year", "2026", "--month", "5",
         "--output-path", str(tmp_path),
         "--start-day", "1", "--end-day", "19"]
    )
    assert rc == 0
    assert (tmp_path / "Schedule_24010_2026-05_01-19.xlsx").exists()
    assert not (tmp_path / "Schedule_24010_2026-05.xlsx").exists()


def test_cli_range_invalid_returns_2(monkeypatch, tmp_path, capsys):
    """end-day past the last day of the month exits cleanly with rc=2."""
    monkeypatch.setattr(cli, "get_member", lambda cid, db: FAKE_MEMBER)
    rc = cli.main(
        ["--center-id", "24010", "--year", "2026", "--month", "4",
         "--output-path", str(tmp_path),
         "--end-day", "31"]  # April has 30 days
    )
    assert rc == 2
    err = capsys.readouterr().err
    assert "end_day 31 exceeds last day of 2026-04" in err


def test_partial_then_full_run_reuses_cached_times(
        monkeypatch, tmp_path):
    """Run May 1-19 with --time-cache, then May 1-31 with the SAME
    cache. Days 1-19 must come back with the exact times they had on
    the partial run; days 20-31 are freshly generated."""
    import openpyxl

    # The fixture stubs load_time_cache/save_time_cache. This test
    # specifically needs the real disk persistence to verify cache
    # round-trip, so put the real implementations back.
    from monthly_schedule.time_cache import (
        load_time_cache as _real_load,
        save_time_cache as _real_save,
    )
    monkeypatch.setattr(cli, "load_time_cache", _real_load)
    monkeypatch.setattr(cli, "save_time_cache", _real_save)
    monkeypatch.setattr(cli, "get_member", lambda cid, db: FAKE_MEMBER)
    cache_path = tmp_path / "tc.json"
    partial_dir = tmp_path / "partial"
    partial_dir.mkdir()
    rc = cli.main(
        ["--center-id", "24010", "--year", "2026", "--month", "5",
         "--output-path", str(partial_dir),
         "--start-day", "1", "--end-day", "19",
         "--time-cache", str(cache_path)]
    )
    assert rc == 0
    partial_wb = openpyxl.load_workbook(
        str(partial_dir / "Schedule_24010_2026-05_01-19.xlsx")
    )
    partial_sheet = partial_wb.active

    # Capture each Date → arrival cell from the Attendance/Trans table.
    # The Pickup time is in the transportation table starting at col F.
    from monthly_schedule.workbook import (
        RIGHT_FIRST_COL, HEADER_ROWS,
    )
    table_header_row = 1 + HEADER_ROWS
    partial_times = {}
    for r in range(table_header_row + 1, partial_sheet.max_row + 1):
        date_cell = partial_sheet.cell(row=r, column=RIGHT_FIRST_COL).value
        pickup = partial_sheet.cell(row=r, column=RIGHT_FIRST_COL + 2).value
        if not isinstance(date_cell, date) or not pickup:
            continue
        partial_times[date_cell] = pickup

    # Re-run for the full month; days 1-19 should match.
    full_dir = tmp_path / "full"
    full_dir.mkdir()
    rc = cli.main(
        ["--center-id", "24010", "--year", "2026", "--month", "5",
         "--output-path", str(full_dir),
         "--time-cache", str(cache_path)]
    )
    assert rc == 0
    full_wb = openpyxl.load_workbook(
        str(full_dir / "Schedule_24010_2026-05.xlsx")
    )
    full_sheet = full_wb.active

    matched = 0
    for r in range(table_header_row + 1, full_sheet.max_row + 1):
        date_cell = full_sheet.cell(row=r, column=RIGHT_FIRST_COL).value
        pickup = full_sheet.cell(row=r, column=RIGHT_FIRST_COL + 2).value
        if not isinstance(date_cell, date):
            continue
        if date_cell in partial_times:
            assert pickup == partial_times[date_cell], (
                f"{date_cell}: full-run pickup {pickup} "
                f"differs from partial-run {partial_times[date_cell]}"
            )
            matched += 1
    assert matched == len(partial_times) > 0


def test_cli_range_feb_29_non_leap_year_returns_2(
        monkeypatch, tmp_path, capsys):
    monkeypatch.setattr(cli, "get_member", lambda cid, db: FAKE_MEMBER)
    rc = cli.main(
        ["--center-id", "24010", "--year", "2026", "--month", "2",
         "--output-path", str(tmp_path),
         "--end-day", "29"]
    )
    assert rc == 2
    assert "exceeds last day" in capsys.readouterr().err


# ---------------------------------------------------------------------------
# --debug flag writes a per-member diagnostic CSV
# ---------------------------------------------------------------------------


def test_debug_flag_writes_per_member_debug_csvs(monkeypatch, tmp_path):
    """--debug produces one Debug_<center_id>_<YYYY-MM>.csv per member
    (no combined file)."""
    monkeypatch.setattr(
        cli, "get_members_by_plan",
        lambda code, db: [FAKE_MEMBER, FAKE_MEMBER_2],
    )
    rc = cli.main(
        ["--plan", "hof", "--year", "2026", "--month", "5",
         "--output-path", str(tmp_path), "--debug"]
    )
    assert rc == 0
    sub = tmp_path / "HOF_2026-05"
    # No combined file.
    assert not (sub / "Debug_2026-05.csv").exists()
    # Debug CSVs live in their own flat folder inside the run's
    # output dir, not next to the timesheets.
    debug_dir = sub / "Debug 2026-05"
    assert not list(sub.glob("Debug_*.csv"))
    for cid in ("24010", "24011"):
        path = debug_dir / f"Debug_{cid}_2026-05.csv"
        assert path.exists()
        lines = path.read_text(encoding="utf-8").splitlines()
        assert lines[0] == (
            "center_id,name,date,day,scheduled,reason,reason_detail,"
            "availability,availability_source,absent,auth_days,"
            "placement_window,max_length,band"
        )
        # Every calendar day of May 2026 gets a row.
        assert len(lines) - 1 == 31
        ids = {line.split(",", 1)[0] for line in lines[1:]}
        assert ids == {cid}


def test_debug_flag_off_does_not_write_debug_csv(monkeypatch, tmp_path):
    monkeypatch.setattr(cli, "get_member", lambda cid, db: FAKE_MEMBER)
    rc = cli.main(
        ["--center-id", "24010", "--year", "2026", "--month", "5",
         "--output-path", str(tmp_path)]
    )
    assert rc == 0
    assert not list(tmp_path.glob("Debug_*.csv"))


def test_write_skipped_members_csv_preserves_input_order(tmp_path):
    """Multiple failures appear in input order."""
    from pathlib import Path

    failures = [
        Failure(24010, "Cheng, Lizhu", "eligibility", "A", None),
        Failure(24011, "Smith, John", "lookup", "B", None),
        Failure(24012, "Jones, Mary", "geocode", "C", None),
    ]
    result = write_skipped_members_csv(
        failures, str(tmp_path), today=date(2026, 6, 4)
    )
    lines = Path(result).read_text(encoding="utf-8").splitlines()
    assert [row.split(",", 1)[0] for row in lines[1:]] == [
        "24010", "24011", "24012",
    ]


def test_collect_debug_rows_uses_travel_adjusted_offsets(monkeypatch):
    import new_monthly_schedule as cli
    from datetime import date
    from monthly_schedule.eligibility_context import MemberContext

    monkeypatch.setattr(cli, "resolve_travel_minutes",
                        lambda member, api_key, cache: 25)
    member = {"center_id": 1, "first_name": "A", "last_name": "B",
              "health_plan": None, "address": "1 Main St",
              "long_lat": "40.0,-73.0"}
    ctx = MemberContext(
        enrollments=[{"id": 1, "center_id": 1,
                      "start_date": date(2026, 1, 1), "end_date": None}],
        authorizations=[{"id": 1, "center_id": 1,
                         "auth_start": date(2026, 1, 1),
                         "auth_end": date(2026, 12, 31),
                         "effective_start": date(2026, 1, 1),
                         "effective_end": date(2026, 12, 31),
                         "auth_days": "1"}],
        absences=[],
        availabilities=[{"id": 1, "center_id": 1,
                         "effective_start_date": date(2026, 1, 1),
                         "effective_end_date": None,
                         "day_of_week": 1,
                         "avail_start": "08:00", "avail_end": "15:00"}],
        one_offs=[],
    )
    rows = cli.collect_debug_rows(member, ctx, 2026, 5,
                                  api_key="K", cache={})
    mon = next(r for r in rows if r["date"] == date(2026, 5, 4))
    # travel 25 + buffer max 5 + drift max 2 = 32 min reserve
    # → out_hi = 15:00 - 32m = 14:28 (default rules, flag on).
    assert mon["placement_window"] == "08:00-14:28"
    assert mon["reason_detail"] == (
        "drop-off reserve 32m before 15:00 avail end"
    )


def test_collect_debug_rows_without_cache_uses_defaults(monkeypatch):
    import new_monthly_schedule as cli
    from datetime import date
    from monthly_schedule.eligibility_context import MemberContext

    def boom(member, api_key, cache):
        raise AssertionError("must not resolve travel without a cache")
    monkeypatch.setattr(cli, "resolve_travel_minutes", boom)
    member = {"center_id": 1, "first_name": "A", "last_name": "B",
              "health_plan": None}
    ctx = MemberContext(
        enrollments=[{"id": 1, "center_id": 1,
                      "start_date": date(2026, 1, 1), "end_date": None}],
        authorizations=[{"id": 1, "center_id": 1,
                         "auth_start": date(2026, 1, 1),
                         "auth_end": date(2026, 12, 31),
                         "effective_start": date(2026, 1, 1),
                         "effective_end": date(2026, 12, 31),
                         "auth_days": "1"}],
        absences=[],
        availabilities=[{"id": 1, "center_id": 1,
                         "effective_start_date": date(2026, 1, 1),
                         "effective_end_date": None,
                         "day_of_week": 1,
                         "avail_start": "08:00", "avail_end": "15:00"}],
        one_offs=[],
    )
    rows = cli.collect_debug_rows(member, ctx, 2026, 5)
    mon = next(r for r in rows if r["date"] == date(2026, 5, 4))
    # Default offsets: drift max 2 + dropoff trail max 12 = 14 min.
    assert mon["placement_window"] == "08:00-14:46"


def test_collect_debug_rows_marks_rows_when_travel_unresolved(monkeypatch):
    import new_monthly_schedule as cli
    from datetime import date
    from monthly_schedule.eligibility_context import MemberContext

    def boom(member, api_key, cache):
        raise cli.TravelError("geocode", "ZERO_RESULTS for 'x'")
    monkeypatch.setattr(cli, "resolve_travel_minutes", boom)
    member = {"center_id": 1, "first_name": "A", "last_name": "B",
              "health_plan": None, "address": "1 Main St",
              "long_lat": "40.0,-73.0"}
    ctx = MemberContext(
        enrollments=[{"id": 1, "center_id": 1,
                      "start_date": date(2026, 1, 1), "end_date": None}],
        authorizations=[{"id": 1, "center_id": 1,
                         "auth_start": date(2026, 1, 1),
                         "auth_end": date(2026, 12, 31),
                         "effective_start": date(2026, 1, 1),
                         "effective_end": date(2026, 12, 31),
                         "auth_days": "1"}],
        absences=[],
        availabilities=[{"id": 1, "center_id": 1,
                         "effective_start_date": date(2026, 1, 1),
                         "effective_end_date": None,
                         "day_of_week": 1,
                         "avail_start": "08:00", "avail_end": "15:00"}],
        one_offs=[],
    )
    rows = cli.collect_debug_rows(member, ctx, 2026, 5,
                                  api_key="K", cache={})
    # Travel unresolved -> falls back to default offsets, and every row
    # is stamped so the CSV reader knows they aren't the real run's.
    mon = next(r for r in rows if r["date"] == date(2026, 5, 4))
    assert mon["placement_window"] == "08:00-14:46"
    assert mon["reason_detail"] == (
        "drop-off reserve 14m before 15:00 avail end; "
        "travel unresolved - default offsets shown"
    )
    # Rows without their own detail still get the stamp.
    assert rows[0]["reason_detail"] == (
        "travel unresolved - default offsets shown"
    )


def test_write_debug_csv_includes_band_column(tmp_path):
    import new_monthly_schedule as cli
    from datetime import date
    rows = [{
        "center_id": 1, "name": "Doe, Jane", "date": date(2026, 5, 4),
        "day": "Mon", "scheduled": True, "reason": "",
        "reason_detail": "", "availability": "",
        "availability_source": "", "absent": "no", "auth_days": "1",
        "placement_window": "08:00-16:00", "max_length": "04:00",
        "band": "morning",
    }]
    path = tmp_path / "debug.csv"
    cli.write_debug_csv(rows, str(path))
    header, data = path.read_text(encoding="utf-8").strip().splitlines()
    assert header.split(",")[-1] == "band"
    assert data.split(",")[-1] == "morning"


def test_collect_debug_rows_carries_band(monkeypatch):
    import new_monthly_schedule as cli
    from datetime import date
    from monthly_schedule.eligibility_context import MemberContext

    member = {"center_id": 1, "first_name": "A", "last_name": "B",
              "health_plan": None}
    ctx = MemberContext(
        enrollments=[{"id": 1, "center_id": 1,
                      "start_date": date(2026, 1, 1), "end_date": None}],
        authorizations=[{"id": 1, "center_id": 1,
                         "auth_start": date(2026, 1, 1),
                         "auth_end": date(2026, 12, 31),
                         "effective_start": date(2026, 1, 1),
                         "effective_end": date(2026, 12, 31),
                         "auth_days": "1"}],
        absences=[], availabilities=[], one_offs=[],
    )
    rows = cli.collect_debug_rows(
        member, ctx, 2026, 5,
        schedule_rules_overrides={"band_enabled": True,
                                  "morning_percent": 100},
    )
    assert rows
    assert all(r["band"] == "morning" for r in rows)


def _on_rows_pipeline(monkeypatch, fail_write=False):
    """Stub process_member's heavy deps: travel resolves instantly,
    build_rows returns a sentinel list, build_workbook writes nothing
    (or raises when fail_write)."""
    import new_monthly_schedule as cli

    sentinel_rows = [{"date": None, "time_in": "09:00"}]
    monkeypatch.setattr(cli, "compute_month_failure",
                        lambda *a, **k: None)
    monkeypatch.setattr(cli, "resolve_travel_minutes",
                        lambda member, api_key, cache: 10)
    monkeypatch.setattr(cli, "build_rows",
                        lambda *a, **k: sentinel_rows)
    if fail_write:
        def boom(*a, **k):
            raise OSError("disk full")
        monkeypatch.setattr(cli, "build_workbook", boom)
    else:
        monkeypatch.setattr(cli, "build_workbook",
                            lambda *a, **k: None)
    return cli, sentinel_rows


def _on_rows_ctx():
    from datetime import date
    from monthly_schedule.eligibility_context import MemberContext

    return MemberContext(
        enrollments=[{"id": 1, "center_id": 1,
                      "start_date": date(2026, 1, 1), "end_date": None}],
        authorizations=[{"id": 1, "center_id": 1,
                         "auth_start": date(2026, 1, 1),
                         "auth_end": date(2026, 12, 31),
                         "effective_start": date(2026, 1, 1),
                         "effective_end": date(2026, 12, 31),
                         "auth_days": "1,3,5"}],
        absences=[], availabilities=[], one_offs=[],
    )


def test_process_member_returns_one_off_conflict_detail(monkeypatch, tmp_path):
    """The structured detail has to reach the caller so the GUI can
    translate the conflict; the CSV keeps the English `reason`."""
    from datetime import date
    import new_monthly_schedule as cli
    from monthly_schedule.eligibility_context import MemberContext

    monkeypatch.setattr(cli, "resolve_travel_minutes",
                        lambda member, api_key, cache: 10)
    ctx = MemberContext(
        enrollments=[{"id": 1, "center_id": 1,
                      "start_date": date(2026, 1, 1), "end_date": None}],
        authorizations=[{"id": 1, "center_id": 1,
                         "auth_start": date(2026, 1, 1),
                         "auth_end": date(2026, 12, 31),
                         "effective_start": date(2026, 1, 1),
                         "effective_end": date(2026, 12, 31),
                         "auth_days": "1,2,3,4,5,6,7"}],
        absences=[{"id": 14, "center_id": 1, "leave_type": "Vacation",
                   "start_date": date(2026, 6, 1),
                   "end_date": date(2026, 6, 2)}],
        availabilities=[],
        one_offs=[{"id": 4, "center_id": 1, "date": date(2026, 6, 1),
                   "avail_start": "09:00", "avail_end": "12:00"}],
    )
    member = {"center_id": 1, "first_name": "A", "last_name": "B",
              "health_plan": "HF"}
    ok, stage, reason, day, detail = cli.process_member(
        member, ctx, 2026, 6, str(tmp_path), "K", {},
    )
    assert ok is False
    assert stage == "one_off_conflict"
    assert day == date(2026, 6, 1)
    assert "OneOffAvailability row 4" in reason
    assert "Absences row 14" in reason
    assert detail["kind"] == "absence"
    assert detail["one_offs"][0]["id"] == 4
    assert detail["absence"]["id"] == 14


def test_process_member_calls_on_rows_after_success(monkeypatch, tmp_path):
    cli, sentinel_rows = _on_rows_pipeline(monkeypatch)
    member = {"center_id": 1, "first_name": "A", "last_name": "B",
              "health_plan": "HF"}
    calls = []
    ok, stage, reason, day, detail = cli.process_member(
        member, _on_rows_ctx(), 2026, 6, str(tmp_path), "K", {},
        on_rows=lambda rows, weekdays: calls.append((rows, weekdays)),
    )
    assert ok is True
    assert calls == [(sentinel_rows, {1, 3, 5})]


def _swap_ctx(authorizations):
    from datetime import date
    from monthly_schedule.eligibility_context import MemberContext

    return MemberContext(
        enrollments=[{"id": 1, "center_id": 1,
                      "start_date": date(2026, 1, 1), "end_date": None}],
        authorizations=authorizations,
        absences=[], availabilities=[], one_offs=[],
    )


def _auth_row(rid, start, end, auth_days):
    return {"id": rid, "center_id": 1,
            "auth_start": start, "auth_end": end,
            "effective_start": start, "effective_end": end,
            "auth_days": auth_days}


def test_process_member_auth_weekdays_use_latest_auth(monkeypatch, tmp_path):
    # Mid-month swap: Mon,Tue for 9/1-9/14, then Wed,Fri from 9/15.
    # The header (and the billing row via on_rows) must show only the
    # later days — {3, 5} — not the whole-month union.
    from datetime import date
    cli, sentinel_rows = _on_rows_pipeline(monkeypatch)
    captured = {}

    def capture_workbook(member, rows, path, auth_weekdays=None):
        captured["weekdays"] = auth_weekdays
    monkeypatch.setattr(cli, "build_workbook", capture_workbook)

    ctx = _swap_ctx([
        _auth_row(1, date(2026, 9, 1), date(2026, 9, 14), "1,2"),
        _auth_row(2, date(2026, 9, 15), date(2026, 9, 30), "3,5"),
    ])
    member = {"center_id": 1, "first_name": "A", "last_name": "B",
              "health_plan": "HF"}
    calls = []
    ok, stage, reason, day, detail = cli.process_member(
        member, ctx, 2026, 9, str(tmp_path), "K", {},
        on_rows=lambda rows, weekdays: calls.append(weekdays),
    )
    assert ok is True
    assert captured["weekdays"] == {3, 5}
    assert calls == [{3, 5}]


def test_process_member_auth_weekdays_from_last_covered_day(
        monkeypatch, tmp_path):
    # Authorization ends mid-month with nothing after: the latest auth
    # in the month is still the one covering 9/14 → its days show.
    from datetime import date
    cli, sentinel_rows = _on_rows_pipeline(monkeypatch)
    captured = {}

    def capture_workbook(member, rows, path, auth_weekdays=None):
        captured["weekdays"] = auth_weekdays
    monkeypatch.setattr(cli, "build_workbook", capture_workbook)

    ctx = _swap_ctx([
        _auth_row(1, date(2026, 9, 1), date(2026, 9, 14), "1,2"),
    ])
    member = {"center_id": 1, "first_name": "A", "last_name": "B",
              "health_plan": "HF"}
    ok, stage, reason, day, detail = cli.process_member(
        member, ctx, 2026, 9, str(tmp_path), "K", {},
    )
    assert ok is True
    assert captured["weekdays"] == {1, 2}


def test_process_member_skips_on_rows_when_write_fails(monkeypatch, tmp_path):
    cli, _rows = _on_rows_pipeline(monkeypatch, fail_write=True)
    member = {"center_id": 1, "first_name": "A", "last_name": "B",
              "health_plan": "HF"}
    calls = []
    ok, stage, reason, day, detail = cli.process_member(
        member, _on_rows_ctx(), 2026, 6, str(tmp_path), "K", {},
        on_rows=lambda rows, weekdays: calls.append(1),
    )
    assert ok is False and stage == "write"
    assert calls == []
