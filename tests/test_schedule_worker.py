import sys
from datetime import date

import pytest

from PyQt6.QtCore import QCoreApplication, QEventLoop


# A QCoreApplication must exist for QThread signal dispatch.
@pytest.fixture(scope="module", autouse=True)
def _qapp():
    app = QCoreApplication.instance() or QCoreApplication(sys.argv)
    yield app


from gui.worker import ScheduleWorker


def _run_to_completion(worker, timeout_ms=5000):
    """Spin a QEventLoop until `worker` emits `finished`. Returns the
    payload tuple (success, payload_dict). A QTimer guard force-quits
    the loop after `timeout_ms` so a regressed worker that never emits
    finished can't hang the test runner forever."""
    from PyQt6.QtCore import QTimer

    loop = QEventLoop()
    captured = []

    def on_finished(success, payload):
        captured.append((success, payload))
        loop.quit()

    timer = QTimer()
    timer.setSingleShot(True)
    timer.timeout.connect(loop.quit)
    timer.start(timeout_ms)

    worker.finished.connect(on_finished)
    worker.start()
    loop.exec()
    timer.stop()
    worker.wait(timeout_ms)
    assert captured, "worker did not emit finished within timeout"
    return captured[0]


def _stub_db_and_caches(monkeypatch, tmp_path):
    """Stub every DB/cache dependency of ScheduleWorker._run_inner so a
    run exercises only the failure-collection and CSV-writing path.
    `get_member` returns None, so the single member fails at lookup."""
    monkeypatch.setattr("gui.worker.get_member", lambda cid, db: None)
    for name in (
        "get_all_enrollments",
        "get_all_authorizations",
        "get_all_absences",
        "get_all_availability",
        "get_all_one_offs",
    ):
        monkeypatch.setattr(f"gui.worker.{name}", lambda db: {})
    monkeypatch.setattr("gui.worker.get_holidays", lambda db: [])
    monkeypatch.setattr(
        "gui.worker.get_operating_days",
        lambda db: [
            {"id": d, "day_name": "", "day_of_week": d,
             "opening_time": "08:00", "closing_time": "16:00"}
            for d in range(1, 8)
        ],
    )
    monkeypatch.setattr("gui.worker.load_cache", lambda path: {})
    monkeypatch.setattr("gui.worker.save_cache", lambda path, cache: None)
    monkeypatch.setattr("gui.worker.load_time_cache", lambda path: {})
    monkeypatch.setattr(
        "gui.worker.save_time_cache", lambda path, cache: None
    )
    monkeypatch.setattr(
        "gui.worker.app_time_cache_path",
        lambda legacy_path=None: str(tmp_path / "time_cache.json"),
    )


def test_worker_payload_carries_one_off_conflict_detail(monkeypatch, tmp_path):
    """The GUI summary translates the conflict from structured detail,
    so the payload must ship it alongside the English reason."""
    _stub_db_and_caches(monkeypatch, tmp_path)
    member = {"center_id": 24010, "last_name": "B", "first_name": "A",
              "health_plan": "HF", "address": "x", "long_lat": "0,0"}
    monkeypatch.setattr("gui.worker.get_member", lambda cid, db: member)

    detail = {
        "kind": "absence",
        "day": "2026-05-04",
        "one_offs": [{"id": 4, "avail_start": "09:00", "avail_end": "12:00"}],
        "absence": {"id": 14, "leave_type": "Vacation",
                    "start_date": "2026-05-04", "end_date": "2026-05-05"},
    }

    def fake_process_member(member, ctx, year, month, out_dir,
                            api_key, cache, on_rows=None, **kwargs):
        return (False, "one_off_conflict", "english sentence",
                date(2026, 5, 4), detail)

    monkeypatch.setattr("gui.worker.process_member", fake_process_member)

    worker = _make_worker(tmp_path)
    success, payload = _run_to_completion(worker)

    assert success is False
    assert len(payload["failures"]) == 1
    failure = payload["failures"][0]
    assert failure["stage"] == "one_off_conflict"
    assert failure["reason"] == "english sentence"
    assert failure["detail"] == detail


def _make_worker(tmp_path):
    return ScheduleWorker(
        mode="single",
        center_id=24010,
        center_ids=None,
        plan_code=None,
        year=2026,
        month=5,
        out_dir=str(tmp_path),
        db_path="unused.accdb",
        google_api_key="K",
        geo_cache=str(tmp_path / "geo_cache.json"),
    )


def test_worker_warns_and_falls_back_when_skipped_csv_locked(
        monkeypatch, tmp_path):
    """When the primary skipped-members CSV is locked (e.g. open in
    Excel), the worker logs a fallback warning and writes the report
    under the `_1` name instead of dying with an unhandled error."""
    _stub_db_and_caches(monkeypatch, tmp_path)
    today = date.today().isoformat()
    (tmp_path / f"skipped_members_{today}.csv").mkdir()

    worker = _make_worker(tmp_path)
    log_events = []
    worker.log_line.connect(lambda key, args: log_events.append((key, args)))

    success, payload = _run_to_completion(worker)

    assert success is False
    assert "error_text" not in payload
    assert payload["failures"][0]["stage"] == "lookup"
    fallback = tmp_path / f"skipped_members_{today}_1.csv"
    assert fallback.exists()
    assert ("worker.skipped_csv_fallback", {
        "primary": f"skipped_members_{today}.csv",
        "filename": f"skipped_members_{today}_1.csv",
    }) in log_events
    assert ("worker.wrote_skipped_csv", {
        "filename": f"skipped_members_{today}_1.csv",
    }) in log_events


def test_worker_warns_when_skipped_csv_unwritable(monkeypatch, tmp_path):
    """When every candidate CSV name is locked, the worker logs a
    warning and still finishes with the normal summary payload instead
    of surfacing an unhandled-error crash."""
    _stub_db_and_caches(monkeypatch, tmp_path)
    today = date.today().isoformat()
    (tmp_path / f"skipped_members_{today}.csv").mkdir()
    for n in range(1, 10):
        (tmp_path / f"skipped_members_{today}_{n}.csv").mkdir()

    worker = _make_worker(tmp_path)
    log_events = []
    worker.log_line.connect(lambda key, args: log_events.append((key, args)))

    success, payload = _run_to_completion(worker)

    assert success is False
    assert "error_text" not in payload
    assert payload["failures"][0]["stage"] == "lookup"
    assert any(
        key == "worker.skipped_csv_failed" for key, _args in log_events
    )
    assert not any(
        key == "worker.wrote_skipped_csv" for key, _args in log_events
    )


def test_worker_single_mode_has_no_billing_workbook(monkeypatch, tmp_path):
    """Non-all modes never build the billing workbook; the payload
    still carries the keys (None) so the summary code stays simple."""
    _stub_db_and_caches(monkeypatch, tmp_path)
    worker = _make_worker(tmp_path)
    success, payload = _run_to_completion(worker)
    assert payload["billing_path"] is None
    assert payload["billing_error"] is None


def test_worker_all_mode_writes_billing_workbook(monkeypatch, tmp_path):
    """An All-Members run collects billing rows via the on_rows hook and
    writes the aggregate billing workbook at the base output dir."""
    from datetime import date as _d

    _stub_db_and_caches(monkeypatch, tmp_path)
    member = {"center_id": 1, "last_name": "B", "first_name": "A",
              "health_plan": "HF", "address": "x", "long_lat": "0,0"}
    monkeypatch.setattr("gui.worker.get_all_members",
                        lambda db: [member])
    monkeypatch.setattr(
        "gui.worker.get_all_billing_fields",
        lambda db: {1: {"gender": "M", "dob": "2/1/1953",
                        "admission_date": "1/1/2024",
                        "medicaid": "AB123"}},
    )

    def fake_process_member(member, ctx, year, month, out_dir,
                            api_key, cache, on_rows=None, **kwargs):
        if on_rows is not None:
            on_rows([{"date": _d(year, month, 1), "time_in": "09:00"}],
                    {1, 3, 5})
        return (True, None, None, None, None)

    monkeypatch.setattr("gui.worker.process_member", fake_process_member)

    worker = _make_worker(tmp_path)
    worker.mode = "all"
    worker.billing_name = "Jane Doe"
    log_events = []
    worker.log_line.connect(lambda key, args: log_events.append(key))

    success, payload = _run_to_completion(worker)

    assert success is True
    assert payload["billing_error"] is None
    assert payload["billing_path"] is not None
    import os as _os
    assert _os.path.exists(payload["billing_path"])
    assert _os.path.dirname(payload["billing_path"]) == str(tmp_path)
    # the configured billing name from Settings drives the filename
    assert _os.path.basename(payload["billing_path"]) == (
        "5. May 2026 billing Jane Doe.xlsx"
    )
    assert "worker.wrote_billing" in log_events

    from openpyxl import load_workbook
    ws = load_workbook(payload["billing_path"])["Sheet1"]
    # the HF member landed in the HealthFirst section with its ID
    found = [c.value for row in ws.iter_rows(min_col=2, max_col=2)
             for c in row]
    assert 1 in found


def test_worker_passes_calendar_to_process_member(monkeypatch, tmp_path):
    """The worker builds the center calendar once per run and hands it
    to process_member, so holidays/closed days reach the generator."""
    _stub_db_and_caches(monkeypatch, tmp_path)
    member = {"center_id": 24010, "last_name": "B", "first_name": "A",
              "health_plan": "HF", "address": "x", "long_lat": "0,0"}
    monkeypatch.setattr("gui.worker.get_member", lambda cid, db: member)
    seen = {}

    def fake_process_member(member, ctx, year, month, out_dir,
                            api_key, cache, **kwargs):
        seen["calendar"] = kwargs.get("calendar")
        return (True, None, None, None, None)

    monkeypatch.setattr("gui.worker.process_member", fake_process_member)

    worker = _make_worker(tmp_path)
    success, payload = _run_to_completion(worker)

    assert success is True
    from monthly_schedule.center_calendar import CenterCalendar
    assert isinstance(seen["calendar"], CenterCalendar)


_ACTIVITIES = {
    f"A{i}": {"name": f"Act {i}", "c_name": f"活动{i}",
              "frequency": "1.2.3.4.5.6.7"}
    for i in range(1, 15)
}


def _fake_process_member_with_rows(year, month):
    from datetime import date as _d

    def fake_process_member(member, ctx, y, m, out_dir,
                            api_key, cache, on_rows=None, **kwargs):
        if on_rows is not None:
            on_rows([{"date": _d(y, m, 1), "time_in": "09:00",
                      "status": "attended"}], {1, 3, 5})
        return (True, None, None, None, None)
    return fake_process_member


def _stub_member_and_activities(monkeypatch, tmp_path, member):
    _stub_db_and_caches(monkeypatch, tmp_path)
    monkeypatch.setattr("gui.worker.get_member", lambda cid, db: member)
    monkeypatch.setattr("gui.worker.get_all_members", lambda db: [member])
    monkeypatch.setattr("gui.worker.get_all_billing_fields",
                        lambda db: {member["center_id"]: {
                            "gender": "F", "dob": "2/1/1953",
                            "admission_date": "1/1/2024",
                            "medicaid": "AB123"}})
    monkeypatch.setattr("gui.worker.get_activities",
                        lambda db: dict(_ACTIVITIES))
    monkeypatch.setattr("gui.worker.process_member",
                        _fake_process_member_with_rows(2026, 5))


def test_worker_bowery_single_mode_writes_activity_log(
        monkeypatch, tmp_path):
    """With a Bowery program name, a single-member run drops the
    activity log next to the timesheet — and keeps it out of the
    Print pipeline (generated_paths)."""
    member = {"center_id": 24010, "last_name": "B", "first_name": "A",
              "health_plan": "HF", "address": "x", "long_lat": "0,0"}
    _stub_member_and_activities(monkeypatch, tmp_path, member)

    worker = _make_worker(tmp_path)
    worker.program_name = "Bowery SADC"
    log_events = []
    worker.log_line.connect(lambda key, args: log_events.append(key))

    success, payload = _run_to_completion(worker)

    assert success is True
    log_path = tmp_path / "(24010) May 2026 Activity log.xlsx"
    assert log_path.exists()
    assert "worker.wrote_activity_log" in log_events
    assert str(log_path) not in payload["generated_paths"]
    assert all("Activity log" not in p for p in payload["generated_paths"])


def test_worker_without_bowery_program_no_activity_log(
        monkeypatch, tmp_path):
    member = {"center_id": 24010, "last_name": "B", "first_name": "A",
              "health_plan": "HF", "address": "x", "long_lat": "0,0"}
    _stub_member_and_activities(monkeypatch, tmp_path, member)
    calls = []
    monkeypatch.setattr("gui.worker.get_activities",
                        lambda db: calls.append(db) or dict(_ACTIVITIES))

    for program in ("", "Chinatown"):
        worker = _make_worker(tmp_path)
        worker.program_name = program
        success, payload = _run_to_completion(worker)
        assert success is True
        assert not (tmp_path / "(24010) May 2026 Activity log.xlsx").exists()
    assert calls == []


def test_worker_all_mode_activity_logs_grouped_by_plan(
        monkeypatch, tmp_path):
    """All-members runs group logs into Activity Logs YYYY-MM/<PLAN>
    even when the MLTC-folders checkbox is off."""
    member = {"center_id": 24010, "last_name": "B", "first_name": "A",
              "health_plan": "HF", "address": "x", "long_lat": "0,0"}
    _stub_member_and_activities(monkeypatch, tmp_path, member)

    worker = _make_worker(tmp_path)
    worker.mode = "all"
    worker.billing_name = "Jane Doe"
    worker.program_name = "bowery"
    assert worker.separate_by_plan is False

    success, payload = _run_to_completion(worker)

    assert success is True
    log_path = (tmp_path / "Activity Logs 2026-05" / "HF"
                / "(24010) May 2026 Activity log.xlsx")
    assert log_path.exists()


@pytest.mark.parametrize("program", ["Bowery", "Cathay"])
def test_worker_activities_table_failure_degrades(
        monkeypatch, tmp_path, program):
    """A missing Activities table warns and skips the logs; the
    timesheets still generate."""
    member = {"center_id": 24010, "last_name": "B", "first_name": "A",
              "health_plan": "HF", "address": "x", "long_lat": "0,0"}
    _stub_member_and_activities(monkeypatch, tmp_path, member)

    def boom(db):
        raise RuntimeError("Could not read the Activities table: nope")
    monkeypatch.setattr("gui.worker.get_activities", boom)

    worker = _make_worker(tmp_path)
    worker.program_name = program
    log_events = []
    worker.log_line.connect(lambda key, args: log_events.append(key))

    success, payload = _run_to_completion(worker)

    assert success is True
    assert "worker.activities_failed" in log_events
    assert not (tmp_path / "(24010) May 2026 Activity log.xlsx").exists()


@pytest.mark.parametrize("program,build_name", [
    ("Bowery", "build_activity_log"),
    ("Cathay", "build_cathay_activity_log"),
])
def test_worker_activity_log_failure_keeps_timesheet(
        monkeypatch, tmp_path, program, build_name):
    """A per-member activity log failure is logged and never fails the
    already-written timesheet."""
    member = {"center_id": 24010, "last_name": "B", "first_name": "A",
              "health_plan": "HF", "address": "x", "long_lat": "0,0"}
    _stub_member_and_activities(monkeypatch, tmp_path, member)

    def boom(*args, **kwargs):
        raise ValueError("template exploded")
    monkeypatch.setattr(f"gui.worker.{build_name}", boom)

    worker = _make_worker(tmp_path)
    worker.program_name = program
    log_events = []
    worker.log_line.connect(lambda key, args: log_events.append((key, args)))

    success, payload = _run_to_completion(worker)

    assert success is True
    assert any(key == "worker.activity_log_failed"
               for key, _args in log_events)


def test_worker_cathay_single_mode_writes_activity_log(
        monkeypatch, tmp_path):
    """A Cathay program name uses the Cathay template: the log lands
    next to the timesheet, carries the auth days in B6, and stays out
    of the Print pipeline."""
    member = {"center_id": 24010, "last_name": "B", "first_name": "A",
              "health_plan": "HF", "address": "x", "long_lat": "0,0"}
    _stub_member_and_activities(monkeypatch, tmp_path, member)

    worker = _make_worker(tmp_path)
    worker.program_name = "Cathay ADC"
    log_events = []
    worker.log_line.connect(lambda key, args: log_events.append(key))

    success, payload = _run_to_completion(worker)

    assert success is True
    log_path = tmp_path / "(24010) May 2026 Activity log.xlsx"
    assert log_path.exists()
    assert "worker.wrote_activity_log" in log_events
    assert all("Activity log" not in p for p in payload["generated_paths"])

    from openpyxl import load_workbook
    ws = load_workbook(str(log_path))["AttndActivityLog"]
    assert "B, A" in ws["B6"].value
    assert "(1.3.5)" in ws["B6"].value  # auth_weekdays reached the builder


def test_worker_cathay_all_mode_groups_by_plan(monkeypatch, tmp_path):
    member = {"center_id": 24010, "last_name": "B", "first_name": "A",
              "health_plan": "HF", "address": "x", "long_lat": "0,0"}
    _stub_member_and_activities(monkeypatch, tmp_path, member)

    worker = _make_worker(tmp_path)
    worker.mode = "all"
    worker.billing_name = "Jane Doe"
    worker.program_name = "cathay"

    success, payload = _run_to_completion(worker)

    assert success is True
    assert (tmp_path / "Activity Logs 2026-05" / "HF"
            / "(24010) May 2026 Activity log.xlsx").exists()


def test_worker_cathay_single_mode_skips_billing_fields(
        monkeypatch, tmp_path):
    """Only the Bowery log needs DOB; a Cathay single-member run must
    not pay for the roster-wide billing-fields query."""
    member = {"center_id": 24010, "last_name": "B", "first_name": "A",
              "health_plan": "HF", "address": "x", "long_lat": "0,0"}
    _stub_member_and_activities(monkeypatch, tmp_path, member)
    calls = []
    monkeypatch.setattr("gui.worker.get_all_billing_fields",
                        lambda db: calls.append(db) or {})

    worker = _make_worker(tmp_path)
    worker.program_name = "Cathay"
    success, _payload = _run_to_completion(worker)
    assert success is True
    assert calls == []

    worker = _make_worker(tmp_path)  # fresh worker, bowery run
    worker.program_name = "Bowery"
    success, _payload = _run_to_completion(worker)
    assert success is True
    assert len(calls) == 1


def test_worker_all_mode_uses_codes_table(monkeypatch, tmp_path):
    """Codes come from db.get_billing_codes; a member whose plan is in
    the table gets its codes, others would get '????'."""
    from datetime import date as _d

    _stub_db_and_caches(monkeypatch, tmp_path)
    member = {"center_id": 2, "last_name": "B", "first_name": "A",
              "health_plan": "HF", "address": "x", "long_lat": "0,0"}
    monkeypatch.setattr("gui.worker.get_all_members",
                        lambda db: [member])
    monkeypatch.setattr("gui.worker.get_all_billing_fields",
                        lambda db: {})
    monkeypatch.setattr("gui.worker.get_billing_codes",
                        lambda db: {"HF": ("S5105", "T2003")})

    def fake_process_member(member, ctx, year, month, out_dir,
                            api_key, cache, on_rows=None, **kwargs):
        if on_rows is not None:
            on_rows([{"date": _d(year, month, 1), "time_in": "09:00"}],
                    {1, 2, 3})
        return (True, None, None, None, None)

    monkeypatch.setattr("gui.worker.process_member", fake_process_member)

    worker = _make_worker(tmp_path)
    worker.mode = "all"
    success, payload = _run_to_completion(worker)

    from openpyxl import load_workbook
    ws = load_workbook(payload["billing_path"])["Sheet1"]
    hf_row = next(r[0].row for r in ws.iter_rows(min_col=2, max_col=2)
                  if r[0].value == 2)
    assert ws.cell(row=hf_row, column=11).value == "S5105"
    assert ws.cell(row=hf_row, column=12).value == "T2003"
